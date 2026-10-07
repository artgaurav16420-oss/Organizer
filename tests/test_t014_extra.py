"""T-014 remainder: cycle e2e, workbook, supersede/move helpers, orphan flow."""
from pathlib import Path

import pytest

import fermi_report_xlsx as fx
from fermi_organizer import fsops
import shutil
from unittest.mock import patch

from fermi_organizer.runmodes import (_detect_supersede_pairs,
                                      _move_children_under_superseding,
                                      run_full)

pytestmark = pytest.mark.skipif(not fx._OPENPYXL_OK,
                                reason="openpyxl not installed")


def test_run_full_cycle_broken_both_placed(tmp_path, make_pdf):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    out = tmp_path / "out"
    make_pdf(in_dir / "F10126106.pdf", [
        "FERMI PART LIST", "F10126107 CHILD PART", "NAME", "Parent A"])
    make_pdf(in_dir / "F10126107.pdf", [
        "FERMI PART LIST", "F10126106 PARENT PART", "NAME", "Child B"])

    ctx = run_full(in_dir, out, False, lambda msg: None, jobs=1)

    assert ctx["counters"]["cycles"] >= 1
    assert ctx["counters"]["copies"] >= 2
    assert (out / "F10126106 Parent A" / "F10126106.pdf").is_file()
    assert (out / "F10126106 Parent A" / "F10126107 Child B" /
            "F10126107.pdf").is_file()


def _minimal_ctx(out):
    return {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 1, "copies": 0,
                      "cycles": 0, "warnings": 0},
        "missing": [],
        "chk": [],
        "orphans": [],
        "roots": [],
        "mismatches": [],
        "used_on_bugs": [],
        "titleblock_mismatches": [],
        "scanned": [],
        "watermarks": [],
        "report_txt": "",
        "names": {},
    }


def test_refresh_from_run_writes_all_sheets(tmp_path):
    out = tmp_path / "out"
    ctx = _minimal_ctx(out)
    xlsx = fx.refresh_from_run(
        output=str(out), run_mode=ctx["run_mode"], run_time=ctx["run_time"],
        counters=ctx["counters"], missing=ctx["missing"], chk=ctx["chk"],
        orphans=ctx["orphans"], report_txt=None, names=ctx["names"],
        roots=ctx["roots"], mismatches=ctx["mismatches"],
        used_on_bugs=ctx["used_on_bugs"],
        titleblock_mismatches=ctx["titleblock_mismatches"],
        scanned=ctx["scanned"], watermarks=ctx["watermarks"])

    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Dashboard", "Missing", "CHK", "Scanned",
                             "Watermark", "Orphans", "Superseded", "Roots",
                             "USED ON check", "Title block check", "Files",
                             "Run History"]
    assert Path(str(xlsx) + ".history.json").is_file()


def test_missing_sheet_builder(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Missing"
    fx._missing_sheet(ws, [("F10126106", "F10126109")],
                      {"F10126109": "Child nine"})
    assert ws["A2"].value == "F10126109"
    assert ws["B2"].value == "Child nine"
    assert ws["C2"].value == "F10126106"


def test_chk_sheet_builder(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "CHK"
    fx._chk_sheet(ws, {"chk": ["F10126107__CHK"]},
                  {"F10126107__CHK": "Bad part"})
    assert ws["A2"].value == "F10126107__CHK"
    assert ws["B2"].value == "Bad part"


@pytest.mark.parametrize("new_stem,org_stem,kind", [
    ("F10126106_A", "F10126106", "pair"),
    ("F10126106", "F10126106_A", "superseded"),
    ("F10126106_A", "F10126106_A", "tie"),
])
def test_detect_supersede_pairs_newer_older_tie(tmp_path, new_stem, org_stem,
                                                kind):
    new_index = {new_stem: tmp_path / f"{new_stem}.pdf"}
    old_new = {}
    organized = {org_stem: [tmp_path / "out" / f"{org_stem}.pdf"]}
    logged = []
    pairs = _detect_supersede_pairs(new_index, old_new, organized,
                                    logged.append)
    if kind == "pair":
        assert pairs == [(new_stem, org_stem)]
        assert new_stem in new_index and not old_new
    elif kind == "superseded":
        assert pairs == []
        assert new_stem not in new_index
        assert old_new == {new_stem: tmp_path / f"{new_stem}.pdf"}
    else:
        assert pairs == []
        assert new_stem in new_index and not old_new
        assert any("ties with organized" in m for m in logged)


@pytest.mark.parametrize("dry_run", [True, False])
def test_move_children_under_superseding(tmp_path, dry_run):
    out = tmp_path / "out"
    s_folder = out / "F10126106 Parent"
    s_folder.mkdir(parents=True)
    (s_folder / "F10126106.pdf").write_bytes(b"%PDF")
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    (child_folder / "F10126107.pdf").write_bytes(b"%PDF")

    organized = {"F10126106": [s_folder / "F10126106.pdf"],
                 "F10126107": [child_folder / "F10126107.pdf"]}
    logged = []
    moved_dirs = {}
    moves = _move_children_under_superseding(
        ["F10126106"], organized, {"F10126106": ["F10126107"]},
        {"F10126106", "F10126107"}, out, dry_run, moved_dirs, logged.append)

    assert moves == 1
    assert moved_dirs == {child_folder: s_folder / "F10126107 Child"}
    if dry_run:
        assert child_folder.is_dir()
        assert not (s_folder / "F10126107 Child").exists()
    else:
        assert not child_folder.exists()
        assert (s_folder / "F10126107 Child" / "F10126107.pdf").is_file()


def test_move_children_under_superseding_oserror(tmp_path):
    out = tmp_path / "out"
    s_folder = out / "F10126106 Parent"
    s_folder.mkdir(parents=True)
    (s_folder / "F10126106.pdf").write_bytes(b"%PDF")
    child_folder1 = out / "F10126107 Child1"
    child_folder1.mkdir(parents=True)
    (child_folder1 / "F10126107.pdf").write_bytes(b"%PDF")
    child_folder2 = out / "F10126108 Child2"
    child_folder2.mkdir(parents=True)
    (child_folder2 / "F10126108.pdf").write_bytes(b"%PDF")

    organized = {"F10126106": [s_folder / "F10126106.pdf"],
                 "F10126107": [child_folder1 / "F10126107.pdf"],
                 "F10126108": [child_folder2 / "F10126108.pdf"]}
    logged = []
    moved_dirs = {}

    orig_move = shutil.move

    def mock_move(src, dst):
        if "F10126107" in src:
            raise OSError("Permission denied")
        return orig_move(src, dst)

    with patch("shutil.move", side_effect=mock_move):
        moves = _move_children_under_superseding(
            ["F10126106"], organized, {"F10126106": ["F10126107", "F10126108"]},
            {"F10126106", "F10126107", "F10126108"}, out, False, moved_dirs, logged.append)

    assert moves == 1
    assert child_folder1 in organized["F10126107"][0].parents
    assert moved_dirs == {child_folder2: s_folder / "F10126108 Child2"}
    text = "\n".join(logged)
    assert "WARNING: move failed for F10126107: Permission denied" in text
    expected_move = (f"moved: {Path('F10126108 Child2')} -> "
                     f"{Path('F10126106 Parent') / 'F10126108 Child2'}")
    assert expected_move in text


def test_copy_orphans_then_retire_adopted(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    src = tmp_path / "F10126108.pdf"
    src.write_bytes(b"%PDF")
    logged = []
    assert fsops.copy_orphans(["F10126108"], {"F10126108": src}, out,
                              False, logged.append) == 1
    orph_path = out / "_orphans" / "F10126108.pdf"
    assert orph_path.is_file()

    live_dir = out / "F10126109 New assy" / "F10126108 Live part"
    live_dir.mkdir(parents=True)
    (live_dir / "F10126108.pdf").write_bytes(b"%PDF")
    retired = fsops.retire_adopted_orphans({"F10126108": orph_path}, out,
                                           False, logged.append)
    assert retired == 1
    assert not orph_path.exists()
    assert "Retired 1 orphan copy(ies)" in "\n".join(logged)
