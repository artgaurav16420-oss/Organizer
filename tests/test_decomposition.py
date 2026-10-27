"""Characterization suite: cycle e2e, workbook, supersede/move helpers, orphan flow."""
import os
import shutil
from pathlib import Path
from unittest.mock import patch

import pytest

import fermi_report_xlsx as fx
from fermi_organizer import fsops, runmodes
from fermi_organizer.runmodes import (
    _adopt_org_children_below,
    _adopt_one_org_child,
    _copy_org_child_tree,
    _detect_supersede_pairs,
    _move_children_under_superseding,
    _place_above_organized_children,
    normalize_folder_names,
    run_full,
)

pytestmark = pytest.mark.skipif(not fx._OPENPYXL_OK,
                                reason="openpyxl not installed")


def test_run_full_cycle_broken_both_placed(tmp_path, make_pdf):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    out = tmp_path / "out"
    make_pdf(in_dir / "F10126106.pdf", [
        "FERMI PART LIST", "F10126107 CHILD PART", "NAME", "Parent A"])
    make_pdf(in_dir / "F10126107.pdf", [
        "FERMI PART LIST", "F10126106 PARENT PART", "NAME", "Child B"])

    ctx = run_full(in_dir, out, False, lambda msg: None, jobs=1)

    assert ctx["counters"]["cycles"] >= 1
    assert ctx["counters"]["copies"] >= 2
    assert (out / "F10126106 Parent A" / "F10126106.pdf").is_file()
    assert (out / "F10126106 Parent A" / "F10126107 Child B" /
            "F10126107.pdf").is_file()


def _minimal_ctx(out):
    return {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 1, "copies": 0,
                      "cycles": 0, "warnings": 0},
        "missing": [],
        "chk": [],
        "orphans": [],
        "roots": [],
        "mismatches": [],
        "used_on_bugs": [],
        "titleblock_mismatches": [],
        "scanned": [],
        "watermarks": [],
        "report_txt": "",
        "names": {},
    }


def test_refresh_from_run_writes_all_sheets(tmp_path):
    out = tmp_path / "out"
    ctx = _minimal_ctx(out)
    xlsx = fx.refresh_from_run(
        output=str(out), run_mode=ctx["run_mode"], run_time=ctx["run_time"],
        counters=ctx["counters"], missing=ctx["missing"], chk=ctx["chk"],
        orphans=ctx["orphans"], report_txt=None, names=ctx["names"],
        roots=ctx["roots"], mismatches=ctx["mismatches"],
        used_on_bugs=ctx["used_on_bugs"],
        titleblock_mismatches=ctx["titleblock_mismatches"],
        scanned=ctx["scanned"], watermarks=ctx["watermarks"])

    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Dashboard", "Missing", "CHK", "Scanned",
                             "Watermark", "Orphans", "Superseded", "Roots",
                             "USED ON check", "Title block check", "Files",
                             "Run History"]
    assert Path(str(xlsx) + ".history.json").is_file()


def test_missing_sheet_builder(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Missing"
    fx._missing_sheet(ws, [("F10126106", "F10126109")],
                      {"F10126109": "Child nine"})
    assert ws["A2"].value == "F10126109"
    assert ws["B2"].value == "Child nine"
    assert ws["C2"].value == "F10126106"


def test_chk_sheet_builder(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "CHK"
    fx._chk_sheet(ws, {"chk": ["F10126107__CHK"]},
                  {"F10126107__CHK": "Bad part"}, lambda stem: None)
    assert ws["A2"].value == "F10126107__CHK"
    assert ws["B2"].value == "Bad part"


@pytest.mark.parametrize("new_stem,org_stem,kind", [
    ("F10126106_A", "F10126106", "pair"),
    ("F10126106", "F10126106_A", "superseded"),
    ("F10126106_A", "F10126106_A", "tie"),
])
def test_detect_supersede_pairs_newer_older_tie(tmp_path, new_stem, org_stem,
                                                kind):
    new_index = {new_stem: tmp_path / f"{new_stem}.pdf"}
    old_new = {}
    organized = {org_stem: [tmp_path / "out" / f"{org_stem}.pdf"]}
    logged = []
    pairs = _detect_supersede_pairs(new_index, old_new, organized,
                                    logged.append)
    if kind == "pair":
        assert pairs == [(new_stem, org_stem)]
        assert new_stem in new_index and not old_new
    elif kind == "superseded":
        assert pairs == []
        assert new_stem not in new_index
        assert old_new == {new_stem: tmp_path / f"{new_stem}.pdf"}
    else:
        assert pairs == []
        assert new_stem in new_index and not old_new
        assert any("ties with organized" in m for m in logged)


@pytest.mark.parametrize("dry_run", [True, False])
def test_move_children_under_superseding(tmp_path, dry_run):
    out = tmp_path / "out"
    s_folder = out / "F10126106 Parent"
    s_folder.mkdir(parents=True)
    (s_folder / "F10126106.pdf").write_bytes(b"%PDF")
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    (child_folder / "F10126107.pdf").write_bytes(b"%PDF")

    organized = {"F10126106": [s_folder / "F10126106.pdf"],
                 "F10126107": [child_folder / "F10126107.pdf"]}
    logged = []
    moved_dirs = {}
    moves, _copies = _move_children_under_superseding(
        ["F10126106"], organized, {"F10126106": ["F10126107"]},
        {"F10126106", "F10126107"}, out, dry_run, moved_dirs, logged.append)

    assert moves == 1
    assert moved_dirs == {child_folder: s_folder / "F10126107 Child"}
    if dry_run:
        assert child_folder.is_dir()
        assert not (s_folder / "F10126107 Child").exists()
    else:
        assert not child_folder.exists()
        assert (s_folder / "F10126107 Child" / "F10126107.pdf").is_file()


@pytest.mark.parametrize("dry_run", [True, False])
def test_move_children_under_superseding_shared_folder_skips(tmp_path, dry_run):
    # A folder holding another drawing's PDF is refused: moving it would take
    # unrelated files along. Covered in dry-run too (the plan must not report
    # a move the real run would refuse).
    out = tmp_path / "out"
    s_folder = out / "F10126106 Parent"
    s_folder.mkdir(parents=True)
    (s_folder / "F10126106.pdf").write_bytes(b"%PDF")
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    (child_folder / "F10126107.pdf").write_bytes(b"%PDF")
    foreign = child_folder / "F10126108.pdf"
    foreign.write_bytes(b"%PDF")

    organized = {"F10126106": [s_folder / "F10126106.pdf"],
                 "F10126107": [child_folder / "F10126107.pdf"],
                 "F10126108": [foreign]}
    logged = []
    moved_dirs = {}
    moves, _copies = _move_children_under_superseding(
        ["F10126106"], organized, {"F10126106": ["F10126107"]},
        {"F10126106", "F10126107"}, out, dry_run, moved_dirs, logged.append)

    assert moves == 0
    assert moved_dirs == {}
    assert child_folder.is_dir()
    assert not (s_folder / "F10126107 Child").exists()
    text = "\n".join(logged)
    assert "F10126107: folder shared with F10126108.pdf" in text
    assert "skipping move" in text


def test_move_children_under_superseding_oserror(tmp_path):
    out = tmp_path / "out"
    s_folder = out / "F10126106 Parent"
    s_folder.mkdir(parents=True)
    (s_folder / "F10126106.pdf").write_bytes(b"%PDF")
    child_folder1 = out / "F10126107 Child1"
    child_folder1.mkdir(parents=True)
    (child_folder1 / "F10126107.pdf").write_bytes(b"%PDF")
    child_folder2 = out / "F10126108 Child2"
    child_folder2.mkdir(parents=True)
    (child_folder2 / "F10126108.pdf").write_bytes(b"%PDF")

    organized = {"F10126106": [s_folder / "F10126106.pdf"],
                 "F10126107": [child_folder1 / "F10126107.pdf"],
                 "F10126108": [child_folder2 / "F10126108.pdf"]}
    logged = []
    moved_dirs = {}

    orig_move = shutil.move

    def mock_move(src, dst):
        if "F10126107" in src:
            raise OSError("Permission denied")
        return orig_move(src, dst)

    with patch("shutil.move", side_effect=mock_move):
        moves, _copies = _move_children_under_superseding(
            ["F10126106"], organized, {"F10126106": ["F10126107", "F10126108"]},
            {"F10126106", "F10126107", "F10126108"}, out, False, moved_dirs, logged.append)

    assert moves == 1
    assert child_folder1 in organized["F10126107"][0].parents
    assert moved_dirs == {child_folder2: s_folder / "F10126108 Child2"}
    text = "\n".join(logged)
    assert "WARNING: move failed for F10126107: Permission denied" in text
    expected_move = (f"moved: {Path('F10126108 Child2')} -> "
                     f"{Path('F10126106 Parent') / 'F10126108 Child2'}")
    assert expected_move in text


@pytest.mark.parametrize("dry_run", [True, False])
def test_place_above_organized_children_shared_folder_skips(tmp_path, dry_run,
                                                            make_pdf):
    out = tmp_path / "out"
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    make_pdf(child_folder / "F10126108.pdf", ["NAME", "Other part"])
    new_pdf = make_pdf(tmp_path / "F10126109.pdf",
                       ["NAME", "Higher assembly"])

    organized = {"F10126107": [child_folder / "F10126107.pdf"],
                 "F10126108": [child_folder / "F10126108.pdf"]}
    new_index = {"F10126109": new_pdf}
    new_children = {"F10126109": set()}
    new_names = {"F10126109": "Higher assembly"}
    moved_dirs = {}
    logged = []

    copies, moves = _place_above_organized_children(
        "F10126109", ["F10126107"], new_children, new_index, new_names,
        organized, moved_dirs, out, dry_run, logged.append,
        org_children_of={})

    assert moves == 0
    assert moved_dirs == {}
    assert child_folder.is_dir()
    assert (child_folder / "F10126108.pdf").is_file()
    target = out / "F10126109 Higher assembly" / "F10126107 Child"
    assert not target.exists()
    text = "\n".join(logged)
    assert "F10126107: folder shared with F10126108.pdf" in text
    assert "skipping move" in text
    if not dry_run:
        # The new root's own PDF is still placed; only the move is refused.
        assert (out / "F10126109 Higher assembly" / "F10126109.pdf").is_file()


def test_nested_child_copy_failure_leaves_no_partial_target(tmp_path, make_pdf,
                                                            monkeypatch):
    out = tmp_path / "out"
    parent = out / "F10126106 Parent"  # live parent: child is nested, so copied
    child_folder = parent / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    new_pdf = make_pdf(tmp_path / "F10126109.pdf", ["NAME", "Higher assembly"])

    def partial_then_fail(src, dst, *a, **k):
        Path(dst).mkdir(parents=True)
        (Path(dst) / "partial.pdf").write_bytes(b"")
        raise OSError("disk full")

    monkeypatch.setattr(shutil, "copytree", partial_then_fail)
    logged = []
    _place_above_organized_children(
        "F10126109", ["F10126107"], {"F10126109": set()},
        {"F10126109": new_pdf}, {"F10126109": "Higher assembly"},
        {"F10126107": [child_folder / "F10126107.pdf"]}, {}, out,
        False, logged.append, org_children_of={})

    target = out / "F10126109 Higher assembly" / "F10126107 Child"
    assert not target.exists()
    assert any("copy failed for F10126107" in m for m in logged)


def test_copy_org_child_tree_copies_shared_folder(tmp_path, make_pdf):
    # The shared-folder guard is move-only: a source folder holding a foreign
    # PDF still lands whole under the target (copying strands nothing).
    out = tmp_path / "out"
    parent = out / "F10126106 Parent"
    child_folder = parent / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    make_pdf(child_folder / "F10126108.pdf", ["NAME", "Other part"])
    organized = {"F10126107": [child_folder / "F10126107.pdf"],
                 "F10126108": [child_folder / "F10126108.pdf"]}
    logged = []
    copies = _copy_org_child_tree("F10126107", out / "NewParent", organized,
                                  {}, out, False, logged.append)
    assert copies == 2
    assert (out / "NewParent" / "F10126107 Child" / "F10126107.pdf").is_file()
    assert child_folder.is_dir()


def test_copy_org_child_tree_skips_overlong_descendant(tmp_path, make_pdf,
                                                       monkeypatch):
    # The folder-level MAX_DIR check does not bound descendant files: a dest
    # path beyond MAX_PATH is skipped up front instead of failing mid-copy.
    # MAX_PATH is shrunk for determinism (tmp-path length varies by machine).
    monkeypatch.setattr("fermi_organizer.runmodes.MAX_PATH", 100)
    out = tmp_path / "out"
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [child_folder / "F10126107.pdf"]}

    calls = []
    monkeypatch.setattr(shutil, "copytree",
                        lambda *a, **k: calls.append((a, k)))
    logged = []
    copies = _copy_org_child_tree("F10126107", out / "NewParent", organized,
                                  {}, out, False, logged.append)
    assert copies == 0
    assert calls == []
    assert any("path too long" in m for m in logged)
    assert not (out / "NewParent").exists()


def test_copy_org_child_tree_non_oserror_cleans_partial_target(tmp_path, make_pdf,
                                                                monkeypatch):
    # Anything past OSError (KeyboardInterrupt, ...) must not leave a partial
    # tree: cleanup runs from finally and the error still propagates.
    # RuntimeError stands in (a real KeyboardInterrupt would abort pytest).
    out = tmp_path / "out"
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [child_folder / "F10126107.pdf"]}

    def blow_up(src, dst, *a, **k):
        Path(dst).mkdir(parents=True)
        (Path(dst) / "partial.pdf").write_bytes(b"")
        raise RuntimeError("interrupted")

    monkeypatch.setattr(shutil, "copytree", blow_up)
    logged = []
    with pytest.raises(RuntimeError, match="interrupted"):
        _copy_org_child_tree("F10126107", out / "NewParent", organized,
                             {}, out, False, logged.append)
    assert not (out / "NewParent" / "F10126107 Child").exists()


def test_copy_orphans_then_retire_adopted(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    src = tmp_path / "F10126108.pdf"
    src.write_bytes(b"%PDF")
    logged = []
    assert fsops.copy_orphans(["F10126108"], {"F10126108": src}, out,
                              False, logged.append) == 1
    orph_path = out / "_orphans" / "F10126108.pdf"
    assert orph_path.is_file()

    live_dir = out / "F10126109 New assy" / "F10126108 Live part"
    live_dir.mkdir(parents=True)
    (live_dir / "F10126108.pdf").write_bytes(b"%PDF")
    retired = fsops.retire_adopted_orphans({"F10126108": orph_path}, out,
                                           False, logged.append)
    assert retired == 1
    assert not orph_path.exists()
    assert "Retired 1 orphan copy(ies)" in "\n".join(logged)
def test_copy_org_child_tree_dry_run_accepts_moved_descendant(tmp_path, make_pdf):
    # A child below a moved ancestor counts as simulated-present in dry-run:
    # exact equality alone would skip a copy the real run makes.
    out = tmp_path / "out"
    new_home = out / "AdoptedRoot"
    current = new_home / "Mid" / "F10126107 Child"
    current.mkdir(parents=True)
    make_pdf(current / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [current / "F10126107.pdf"]}
    moved_dirs = {out / "OldRoot": new_home}
    logged = []
    copies = _copy_org_child_tree("F10126107", out / "Target", organized,
                                  moved_dirs, out, True, logged.append)
    assert copies == 1
    assert any("[DRY-RUN] copy" in m for m in logged)


def test_copy_org_child_tree_counts_relocated_source(tmp_path, make_pdf):
    # After an ancestor move, the copy reads from the relocated folder
    # while organized still indexes the old path: the count must follow
    # the indexed tally, not drop to zero.
    out = tmp_path / "out"
    old_home = out / "OldRoot" / "F10126107 Child"
    old_home.mkdir(parents=True)
    make_pdf(old_home / "F10126107.pdf", ["NAME", "Child part"])
    new_home = out / "NewHome" / "F10126107 Child"
    new_home.mkdir(parents=True)
    make_pdf(new_home / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [old_home / "F10126107.pdf"]}
    moved_dirs = {out / "OldRoot": out / "NewHome"}
    logged = []
    copies = _copy_org_child_tree("F10126107", out / "Target", organized,
                                  moved_dirs, out, False, logged.append)
    assert copies == 1
    assert (out / "Target" / "F10126107 Child" / "F10126107.pdf").is_file()


def test_copy_org_child_tree_lock_loss_stops_before_copy(tmp_path, make_pdf,
                                                          monkeypatch):
    # A steal racing placement must stop the adoption like it stops
    # place_files (LockLostError), not write into the other run's tree.
    from fermi_organizer import runmodes as _rm
    out = tmp_path / "out"
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [child_folder / "F10126107.pdf"]}
    logged = []

    def _boom(lock, log):
        raise fsops.LockLostError("stolen")

    monkeypatch.setattr(_rm, "_ensure_lock", _boom)
    with pytest.raises(fsops.LockLostError):
        _copy_org_child_tree("F10126107", out / "NewParent", organized,
                             {}, out, False, logged.append)
    assert not (out / "NewParent").exists()


def test_copy_org_child_tree_refuses_an_unmeasurable_source(tmp_path, make_pdf,
                                                            monkeypatch):
    # None means "cannot list the subtree", not "short enough". Skipping the
    # guard would copy past MAX_PATH unguarded and fail half-way.
    out = tmp_path / "out"
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [child_folder / "F10126107.pdf"]}
    monkeypatch.setattr("fermi_organizer.runmodes.longest_dest_path",
                        lambda target, current, cache=None: None)

    logged = []
    copies = _copy_org_child_tree("F10126107", out / "NewParent", organized,
                                  {}, out, False, logged.append)
    assert copies == 0
    assert any("cannot measure" in m for m in logged)
    assert not (out / "NewParent").exists()


def test_copy_org_child_tree_measures_the_live_source(tmp_path, make_pdf,
                                                      monkeypatch):
    # After a simulated move in a dry run, `current` is virtual. The subtree
    # must be measured where it still sits on disk, or the copy is planned
    # against a path no scanner can see (and therefore unbounded).
    out = tmp_path / "out"
    live = out / "A"
    live.mkdir(parents=True)
    make_pdf(live / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [out / "B" / "F10126107.pdf"]}
    moved_dirs = {out / "A": out / "B"}
    seen = []
    real = runmodes.longest_dest_path

    def spy(target, current, cache=None):
        seen.append(current)
        return real(target, current, cache)

    monkeypatch.setattr(runmodes, "longest_dest_path", spy)
    logged = []
    copies = _copy_org_child_tree("F10126107", out / "Target", organized,
                                  moved_dirs, out, True, logged.append)
    assert copies == 1
    assert seen == [out / "A"], seen


def test_adopt_one_org_child_refuses_an_unmeasurable_move(tmp_path, make_pdf,
                                                          monkeypatch):
    # Same rule on the re-home path: an unknown descendant length must refuse
    # the move, not skip the guard and move anyway.
    out = tmp_path / "out"
    top = out / "F10126107 Child"
    top.mkdir(parents=True)
    make_pdf(top / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [top / "F10126107.pdf"]}
    monkeypatch.setattr("fermi_organizer.runmodes.longest_dest_path",
                        lambda target, current, cache=None: None)

    logged = []
    copies, moves = _adopt_one_org_child(
        "F10126107", out / "NewParent", organized, {}, out, True,
        logged.append)
    assert (copies, moves) == (0, 0)
    assert any("cannot measure" in m for m in logged)
    assert top.is_dir()


@pytest.mark.skipif(os.name != "nt", reason="UTF-16 accounting is Windows-specific")
def test_copy_guard_counts_utf16_units(tmp_path, make_pdf, monkeypatch):
    # A supplementary-plane character is one code point but two UTF-16 units.
    # The limit is placed between the two counts: a code-point guard copies to
    # a path the OS rejects, a unit-aware one refuses it.
    out = tmp_path / "out"
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [child_folder / "F10126107.pdf"]}

    target = out / ("\U0001F600" * 30)
    target_path = target / "F10126107 Child"
    deepest = target_path / "F10126107.pdf"
    # The copy is staged under a longer sibling name first, so a code-point
    # guard already trips on that suffix. Place the limit above it and below the
    # UTF-16 count, so only unit-aware accounting can refuse this copy.
    cp = len(str(deepest))
    units = len(str(deepest).encode("utf-16-le")) // 2
    staging_extra = (len(f"{target_path.name}.copy_tmp.{os.getpid()}")
                     - len(target_path.name))
    limit = cp + staging_extra + 1
    monkeypatch.setattr("fermi_organizer.runmodes.MAX_PATH", limit)
    assert limit < units

    logged = []
    copies = _copy_org_child_tree("F10126107", target, organized, {}, out,
                                  False, logged.append)
    assert copies == 0
    assert any("path too long" in m for m in logged)


def test_normalize_walks_each_directory_once(tmp_path, make_pdf, monkeypatch):
    # A sheet set adds a candidate per sheet and a refused rename is retried for
    # the next one; nested folders overlap those subtrees. Each directory must
    # be listed once for the whole pass, or normalization is quadratic.
    out = tmp_path / "out"
    top = out / "OLD"
    sub = top / "SUB"
    sub.mkdir(parents=True)
    paths = []
    for sheet in ("", "___DWG1", "___DWG2"):
        path = sub / f"F10126106{sheet}.pdf"
        make_pdf(path, ["NAME", "Part"])
        paths.append(path)
    make_pdf(top / "F10126107.pdf", ["NAME", "Top part"])
    organized = {"F10126106": paths, "F10126107": [top / "F10126107.pdf"]}
    names = {"F10126106": "X" * 200, "F10126107": "Y" * 200}

    seen = []
    real = os.scandir

    def spy(path):
        seen.append(str(path))
        return real(path)

    monkeypatch.setattr(os, "scandir", spy)
    logged = []
    renamed = normalize_folder_names(organized, names, out, True, logged.append,
                                     {})
    assert renamed == 0
    assert any("path too long" in m for m in logged)
    assert len(seen) == len(set(seen)), seen
    assert len(set(seen)) == 2, seen


def test_copy_org_child_tree_leaves_a_preexisting_staging_tree(tmp_path, make_pdf):
    # A PID can be reused and a crashed run can leave a partial tree at the
    # staging name. A failed copytree must not delete it: it is not ours.
    out = tmp_path / "out"
    child_folder = out / "F10126107 Child"
    child_folder.mkdir(parents=True)
    make_pdf(child_folder / "F10126107.pdf", ["NAME", "Child part"])
    organized = {"F10126107": [child_folder / "F10126107.pdf"]}
    target_folder = out / "NewParent"
    target_folder.mkdir()
    staging = target_folder / f"F10126107 Child.copy_tmp.{os.getpid()}"
    staging.mkdir()
    marker = staging / "from_a_crashed_run.pdf"
    marker.write_bytes(b"%PDF-1.4\nkeep")

    logged = []
    copies = _copy_org_child_tree("F10126107", target_folder, organized, {},
                                  out, False, logged.append)
    assert copies == 0
    assert marker.is_file(), "a staging tree this run did not create was deleted"
    assert any("staging path already present" in m for m in logged)
    assert not (target_folder / "F10126107 Child").exists()


def test_superseding_copy_keeps_the_tuple_contract(tmp_path, make_pdf,
                                                   monkeypatch):
    # This branch returns (moves, copies). An early "return 0" made the caller
    # unpack an int and abort repair with TypeError instead of skipping one
    # child.
    out = tmp_path / "out"
    parent = out / "F10126106 Parent"
    parent.mkdir(parents=True)
    make_pdf(parent / "F10126106.pdf", ["NAME", "Parent"])
    first = out / "F10126107 A"
    first.mkdir()
    make_pdf(first / "F10126107.pdf", ["NAME", "Child"])
    second = out / "F10126107 B"
    second.mkdir()
    make_pdf(second / "F10126107.pdf", ["NAME", "Child"])
    organized = {"F10126106": [parent / "F10126106.pdf"],
                 "F10126107": [first / "F10126107.pdf",
                               second / "F10126107.pdf"]}
    monkeypatch.setattr(shutil, "copytree",
                        lambda *a, **k: (_ for _ in ()).throw(FileExistsError()))

    moves, copies = _move_children_under_superseding(
        ["F10126106"], organized, {"F10126106": ["F10126107"]},
        ["F10126106", "F10126107"], out, False, {}, lambda m: None)
    assert (moves, copies) == (0, 0)


def test_organized_child_adoption_counts_utf16_units(tmp_path, monkeypatch):
    # MAX_DIR is a Windows limit too, so the folder guard must count UTF-16
    # units: 20 astral characters are 20 code points and 40 units, and the
    # limit placed between the two must still refuse the folder.
    out = tmp_path / "out"
    folder = out / ("\U0001F600" * 20)
    folder.mkdir(parents=True)
    code_points = len(str(folder))
    units = len(str(folder).encode("utf-16-le")) // 2
    monkeypatch.setattr("fermi_organizer.runmodes.MAX_DIR", code_points + 5)
    assert code_points + 5 < units

    logged = []
    copies, moves = _adopt_org_children_below(
        {"F10126106": folder}, {}, {}, {}, {}, {}, {}, out, True, logged.append)
    assert (copies, moves) == (0, 0)
    assert any("path too long" in m for m in logged)

