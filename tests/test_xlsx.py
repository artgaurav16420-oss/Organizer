"""Workbook builder + history sidecar tests."""
import json
from pathlib import Path

import pytest

import fermi_report_xlsx as fx

pytestmark = pytest.mark.skipif(not fx._OPENPYXL_OK,
                                reason="openpyxl not installed")


def test_load_history_corrupt_json_preserves_corrupt_copy(tmp_path, capsys):
    xlsx = tmp_path / "organize_fermi_report.xlsx"
    sidecar = Path(str(xlsx) + ".history.json")
    sidecar.write_text("{not valid json", encoding="utf-8")

    assert fx._load_history(xlsx) == []
    assert Path(str(sidecar) + ".corrupt").exists()
    assert sidecar.exists()
    out = capsys.readouterr().out
    assert "run history not loaded" in out
    # Warning names the exception type and includes the full traceback.
    assert "JSONDecodeError" in out
    assert "Traceback (most recent call last)" in out


def test_load_history_copy_oserror_handled_gracefully(tmp_path, monkeypatch, capsys):
    xlsx = tmp_path / "organize_fermi_report.xlsx"
    sidecar = Path(str(xlsx) + ".history.json")
    sidecar.write_text("{not valid json", encoding="utf-8")

    def mock_copy2(src, dst):
        raise OSError("Permission denied")

    monkeypatch.setattr(fx.shutil, "copy2", mock_copy2)

    assert fx._load_history(xlsx) == []
    assert not Path(str(sidecar) + ".corrupt").exists()
    assert "run history not loaded" in capsys.readouterr().out


def test_load_history_missing_sidecar_returns_empty_quietly(tmp_path, capsys):
    assert fx._load_history(tmp_path / "nope.xlsx") == []
    assert "run history not loaded" not in capsys.readouterr().out


@pytest.mark.parametrize("payload", ["{}", "123", '"string_data"', "true", "45.67"])
def test_load_history_non_list_json_preserves_corrupt_copy(tmp_path, payload):
    xlsx = tmp_path / "organize_fermi_report.xlsx"
    sidecar = Path(str(xlsx) + ".history.json")
    sidecar.write_text(payload, encoding="utf-8")
    logged = []

    assert fx._load_history(xlsx, logged.append) == []
    assert Path(str(sidecar) + ".corrupt").exists()
    assert "run history is not a list" in "\n".join(logged)


def test_save_history_oserror_warns_without_crash(tmp_path, monkeypatch):
    xlsx = tmp_path / "organize_fermi_report.xlsx"
    logged = []

    def boom(src, dst):
        raise OSError("locked")

    monkeypatch.setattr(fx.os, "replace", boom)
    fx._save_history(xlsx, [["run", 1]], logged.append)

    assert "run history not saved" in "\n".join(logged)


def test_save_history_roundtrip_and_cap(tmp_path):
    xlsx = tmp_path / "organize_fermi_report.xlsx"
    rows = [["run", i] for i in range(250)]

    fx._save_history(xlsx, rows)

    loaded = fx._load_history(xlsx)
    assert len(loaded) == 200
    assert loaded == rows[-200:]
    assert not Path(str(xlsx) + ".history.json.tmp").exists()


def test_build_workbook_creates_expected_sheets(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 1, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [],
        "chk": [],
        "orphans": [],
        "roots": [],
        "mismatches": [],
        "report_txt": "",
        "names": {},
    }

    result = fx.build_workbook(path, ctx)

    assert result == path
    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb.sheetnames == ["Dashboard", "Missing", "CHK", "Scanned",
                             "Watermark", "Orphans", "Superseded", "Roots",
                             "USED ON check", "Title block check", "Files",
                             "Run History"]
    history = json.loads(Path(str(path) + ".history.json").read_text(encoding="utf-8"))
    assert len(history) == 1
    assert history[0][0] == "2026-09-30T12:00:00"


def test_build_workbook_sheet_cells(tmp_path):
    out = tmp_path / "out"
    (out / "F10126106 Assembly").mkdir(parents=True)
    (out / "F10126106 Assembly" / "F10126106.pdf").write_bytes(b"%PDF")
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 3, "roots": 1, "copies": 1,
                     "cycles": 0, "warnings": 0},
        "missing": [("F10126106", "F10126109")],
        "chk": ["F10126107__CHK"],
        "orphans": ["F10126108"],
        "roots": [("F10126106", ["F10126100"])],
        "mismatches": [("F10126106", "F10126107", ["F10126100"])],
        "titleblock_mismatches": [("F10126107__CHK", "revision", "-", "B")],
        "report_txt": "report.txt",
        "names": {"F10126109": "Child nine"},
    }

    fx.build_workbook(path, ctx)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb["Dashboard"]["B5"].value == "PDF FILES IN TREE (LIVE)"
    assert wb["Dashboard"]["B6"].value == 1
    assert wb["Missing"]["A2"].value == "F10126109"
    assert wb["Missing"]["B2"].value == "Child nine"
    assert wb["CHK"]["A2"].value == "F10126107__CHK"
    assert wb["Orphans"]["A2"].value == "F10126108"
    assert wb["Roots"]["C2"].value == "F10126100"
    assert wb["USED ON check"]["E2"].value == "F10126100"
    assert wb["Title block check"]["A2"].value == "F10126107__CHK"
    assert wb["Title block check"]["C2"].value == "revision"
    assert wb["Title block check"]["D2"].value == "-"
    assert wb["Title block check"]["E2"].value == "B"
    assert wb["Files"]["A2"].value == "F10126106"
    hist = wb["Run History"]
    assert hist["A2"].value == "2026-09-30T12:00:00"
    assert hist["C2"].value == 3
    assert hist["E2"].value == 1
    assert hist["H2"].value == "report.txt"


def test_files_sheet_lists_tree_only_not_orphans(tmp_path):
    # The Files sheet is the organized tree; parked _orphans/ files belong on
    # the Orphans sheet and must not be repeated here.
    out = tmp_path / "out"
    (out / "F10126106 Assembly").mkdir(parents=True)
    (out / "F10126106 Assembly" / "F10126106.pdf").write_bytes(b"%PDF")
    orph = out / "_orphans"
    orph.mkdir()
    (orph / "F10126108.pdf").write_bytes(b"%PDF")
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 2, "roots": 1, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [], "chk": [], "orphans": ["F10126108"], "roots": [],
        "mismatches": [], "report_txt": "", "names": {},
    }

    fx.build_workbook(path, ctx)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    files = wb["Files"]
    stems = [v for v in (files.cell(row=r, column=1).value
                         for r in range(2, files.max_row + 1)) if v]
    assert stems == ["F10126106"]
    statuses = [files.cell(row=r, column=4).value
                for r in range(2, files.max_row + 1)]
    assert "ORPHAN (parked)" not in statuses
    # The parked orphan is still reported on its own sheet.
    assert wb["Orphans"]["A2"].value == "F10126108"


def test_file_link_is_relative_and_fits_hyperlink_limit(tmp_path):
    # _file_link is pure path math: no need to create the deep tree on disk.
    out = tmp_path / "out"
    deep = out
    for part in ("F10119081 SSR2 CAVITY ASSY", "F10118801 WELDMENT, SSR2 JACKETED CAVITY",
                 "F10118539 WELDMENT, SSR2 CAVITY", "F10119290 WLDMNT, ENDWALL SUB-ASSY 1",
                 "F10118615 INNER COUPLING RING"):
        deep = deep / part
    pdf = deep / "F10118615_A_DWG1.pdf"

    rel = fx._file_link(pdf, out)
    assert rel == ("F10119081 SSR2 CAVITY ASSY/F10118801 WELDMENT, SSR2 JACKETED CAVITY/"
                   "F10118539 WELDMENT, SSR2 CAVITY/F10119290 WLDMNT, ENDWALL SUB-ASSY 1/"
                   "F10118615 INNER COUPLING RING/F10118615_A_DWG1.pdf")
    assert not rel.startswith("file:")
    assert len(rel) < 255
    # A file outside the output folder falls back to the absolute URI.
    outside = tmp_path / "elsewhere.pdf"
    assert fx._file_link(outside, out).startswith("file://")


def test_files_sheet_links_deep_tree_rows(tmp_path):
    # Deep space/comma-heavy rows must still be clickable: a relative target
    # stays under Excel's 255-char hyperlink limit where a file:/// URI
    # (percent-encoded) does not. Short on-disk names keep tmp_path under
    # Windows MAX_PATH; the link is what this checks.
    out = tmp_path / "out"
    d = out / "A B SSR2 CAVITY ASSY" / "C, D WELDMENT" / "E, F INNER RING"
    d.mkdir(parents=True)
    (d / "F10126106_A_DWG1.pdf").write_bytes(b"%PDF")
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 1, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [], "chk": [], "orphans": [], "roots": [],
        "mismatches": [], "report_txt": "", "names": {},
    }

    fx.build_workbook(path, ctx)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    files = wb["Files"]
    linked = [files.cell(row=r, column=1)
              for r in range(2, files.max_row + 1) if files.cell(row=r, column=1).value]
    assert linked
    assert all(c.hyperlink is not None for c in linked)
    assert linked[0].hyperlink.target.endswith("F10126106_A_DWG1.pdf")
    assert not linked[0].hyperlink.target.startswith("file:")


def test_build_workbook_scanned_and_watermark_sheets(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 2, "roots": 0, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [],
        "chk": [],
        "orphans": [],
        "roots": [],
        "mismatches": [],
        "scanned": [("F10126108", [1])],
        "watermarks": [("F10126109", "watermark text: PRELIMINARY")],
        "report_txt": "",
        "names": {"F10126108": "Loose part", "F10126109": "Stamped part"},
    }

    fx.build_workbook(path, ctx)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    scanned = wb["Scanned"]
    assert scanned["A2"].value == "F10126108"
    assert scanned["B2"].value == "Loose part"
    assert scanned["C2"].value == "1"
    wm = wb["Watermark"]
    assert wm["A2"].value == "F10126109"
    assert wm["B2"].value == "Stamped part"
    assert wm["C2"].value == "watermark text: PRELIMINARY"


def test_build_workbook_lists_used_on_bugs(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 1, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [],
        "chk": [],
        "orphans": [],
        "roots": [],
        "mismatches": [],
        "used_on_bugs": [("F10126108", "F10126107")],
        "report_txt": "",
        "names": {},
    }

    fx.build_workbook(path, ctx)

    from openpyxl import load_workbook
    ws = load_workbook(path)["USED ON check"]
    col_a = [c.value for c in ws["A"]]
    col_c = [c.value for c in ws["C"]]
    assert "USED ON BUGS (1)" in col_a
    assert "F10126108" in col_a
    assert "F10126107" in col_c


def test_resolve_names_exact_match_and_fallback():
    live = {
        "tree": [Path("F10126106_REV_A.pdf"), Path("F10126107.pdf")],
        "orph": [Path("F10126108_B.pdf")],
    }
    names_raw = {
        "F10126106": "Base Bracket",
        "F10126107": "Direct Name",
        "F10126108_B": "Specific Orphan Rev",
        "F10126109": "Standalone Extra",
    }

    resolved = fx._resolve_names(names_raw, live)

    assert resolved == {
        "F10126106_REV_A": "Base Bracket",
        "F10126106": "Base Bracket",
        "F10126107": "Direct Name",
        "F10126108_B": "Specific Orphan Rev",
        "F10126109": "Standalone Extra",
    }


def test_resolve_names_exact_match_takes_precedence():
    live = {
        "tree": [Path("F10126106_REV_A.pdf")],
        "orph": [],
    }
    names_raw = {
        "F10126106": "Base Bracket",
        "F10126106_REV_A": "Specific Rev Bracket",
    }

    resolved = fx._resolve_names(names_raw, live)

    assert resolved["F10126106_REV_A"] == "Specific Rev Bracket"
    assert resolved["F10126106"] == "Base Bracket"


def test_resolve_names_falsy_value_triggers_fallback():
    live = {
        "tree": [Path("F10126106_REV_A.pdf")],
        "orph": [],
    }
    names_raw = {
        "F10126106_REV_A": "",
        "F10126106": "Base Bracket",
    }

    resolved = fx._resolve_names(names_raw, live)

    assert resolved["F10126106_REV_A"] == "Base Bracket"


def test_resolve_names_unresolved_and_empty_inputs():
    live = {
        "tree": [Path("F10126106.pdf"), Path("invalid_file.txt")],
        "orph": [],
    }
    names_raw = {}

    resolved = fx._resolve_names(names_raw, live)
    assert resolved == {}

    empty_resolved = fx._resolve_names({}, {"tree": [], "orph": []})
    assert empty_resolved == {}


def test_resolve_names_chk_reuses_dwg_twin():
    # Superseded CHK stems are never scanned: same base+rev reuses the DWG
    # twin's NAME instead of showing '-'.
    live = {"tree": [], "orph": []}
    names_raw = {"F10137247_B_DWG1": "TUBE, INTERFACE"}
    resolved = fx._resolve_names(names_raw, live,
                                 extra=["F10137247_B_CHK"])
    assert resolved["F10137247_B_CHK"] == "TUBE, INTERFACE"


def test_resolve_names_blocks_cross_sheet_fallback():
    # Sibling sheets have their own content: DWG2 never inherits DWG1's NAME.
    live = {"tree": [], "orph": []}
    names_raw = {"F10126106_A___DWG1": "Sheet one"}
    resolved = fx._resolve_names(names_raw, live,
                                 extra=["F10126106_A___DWG2"])
    assert "F10126106_A___DWG2" not in resolved


def test_resolve_names_sheetless_plain_blocked_from_sheet():
    # A plain stem is its own sheet: it must not inherit a sibling sheet's
    # NAME (either direction) - only CHK requesters bypass sheet matching.
    live = {"tree": [], "orph": []}
    names_raw = {"F10126106_A___DWG1": "Sheet one"}
    resolved = fx._resolve_names(names_raw, live,
                                 extra=["F10126106",
                                        "F10126106_A___DWG2"])
    assert "F10126106" not in resolved
    assert "F10126106_A___DWG2" not in resolved


def test_fallback_name_requires_exact_rev_match():
    # An unrevisioned stem must not inherit a later revision's NAME, and
    # mismatched revisions never mix - only CHK requesters bypass the match.
    live = {"tree": [], "orph": []}
    resolved = fx._resolve_names({"F10126106_B": "Rev B name"}, live,
                                 extra=["F10126106"])
    assert "F10126106" not in resolved
    resolved = fx._resolve_names({"F10126106_A": "Rev A name"}, live,
                                 extra=["F10126106_B"])
    assert "F10126106_B" not in resolved


def test_fallback_name_chk_requester_exempt_from_rev_match():
    # A CHK is the same drawing as its rev twin: an unrevisioned CHK may
    # reuse a revved donor's NAME.
    live = {"tree": [], "orph": []}
    resolved = fx._resolve_names({"F10126106_B": "Rev B name"}, live,
                                 extra=["F10126106_CHK"])
    assert resolved["F10126106_CHK"] == "Rev B name"


def test_resolve_names_bare_base_key_not_leaked_to_sheeted_sibling():
    # A prior-run bare-stem NAME must not leak to sheeted siblings via the
    # base-key lookup: each sheet has its own NAME.
    live = {"tree": [], "orph": []}
    resolved = fx._resolve_names({"F10126106": "Base name"}, live,
                                 extra=["F10126106_A___DWG2"])
    assert resolved["F10126106"] == "Base name"
    assert "F10126106_A___DWG2" not in resolved


def test_resolve_names_sheeted_chk_blocked_from_bare_base_key():
    # A sheeted CHK is still that sheet's content: a bare (sheetless) base
    # NAME must not leak to it via the base-key lookup - same sheet rule as
    # _fallback_name. Only a sheetless CHK may reuse the bare base key.
    live = {"tree": [], "orph": []}
    resolved = fx._resolve_names({"F10126106": "Base name"}, live,
                                 extra=["F10126106_A___DWG2_CHK"])
    assert "F10126106_A___DWG2_CHK" not in resolved
    resolved = fx._resolve_names({"F10126106": "Base name"}, live,
                                 extra=["F10126106_CHK"])
    assert resolved["F10126106_CHK"] == "Base name"


def test_fallback_name_sheeted_chk_requires_same_sheet():
    # DWG2_CHK must not inherit a DWG1 donor's title; it reuses a same-sheet
    # donor (CHK is still exempt from revision matching). A sheetless CHK
    # keeps the cross-sheet carve-out (same drawing as its twin).
    live = {"tree": [], "orph": []}
    resolved = fx._resolve_names({"F10126106_A___DWG1": "Sheet one"}, live,
                                 extra=["F10126106_A___DWG2_CHK"])
    assert "F10126106_A___DWG2_CHK" not in resolved
    resolved = fx._resolve_names({"F10126106_A___DWG1": "Sheet one",
                                 "F10126106_B___DWG2": "Sheet two rev B"},
                                live, extra=["F10126106_A___DWG2_CHK"])
    assert resolved["F10126106_A___DWG2_CHK"] == "Sheet two rev B"
    resolved = fx._resolve_names({"F10137247_B_DWG1": "TUBE, INTERFACE"},
                                live, extra=["F10137247_B_CHK"])
    assert resolved["F10137247_B_CHK"] == "TUBE, INTERFACE"


def test_resolve_names_base_key_requires_rev_equality():
    # A bare-base NAME carries no revision: a revved sheetless stem must not
    # inherit it (would label rev B with another revision's title). Only an
    # unrevisioned sheetless stem - or a CHK, the same drawing as its twin -
    # may reuse the bare base key.
    live = {"tree": [], "orph": []}
    resolved = fx._resolve_names({"F10126106": "Base name"}, live,
                                 extra=["F10126106_B"])
    assert "F10126106_B" not in resolved
    resolved = fx._resolve_names({"F10126106": "Base name"}, live,
                                 extra=["F10126106_B_CHK"])
    assert resolved["F10126106_B_CHK"] == "Base name"


def test_fallback_name_revisionless_chk_prefers_newest_rev():
    # A revisionless CHK is exempt from rev matching, but all revisioned
    # donors previously tied at the same priority: lexical order picked rev
    # A over newer B (stale NAME). Newest revision rank wins now.
    live = {"tree": [], "orph": []}
    resolved = fx._resolve_names({"F10126106_A": "Rev A name",
                                 "F10126106_B": "Rev B name"},
                                live, extra=["F10126106_CHK"])
    assert resolved["F10126106_CHK"] == "Rev B name"


def test_resolve_names_extracted_name_always_wins():
    live = {"tree": [], "orph": []}
    names_raw = {"F10126106_A___DWG1": "Sheet one",
                 "F10126106_A___DWG2": "Sheet two"}
    resolved = fx._resolve_names(names_raw, live,
                                 extra=["F10126106_A___DWG2"])
    assert resolved["F10126106_A___DWG2"] == "Sheet two"


def test_prior_names_reads_files_sheet(tmp_path):
    # Names established by earlier runs persist: Files stems map back,
    # '-' placeholders and junk rows are skipped.
    from openpyxl import Workbook
    out = tmp_path / "organize_fermi_report.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Files"
    ws.append(["Drawing (stem)", "Name"])
    ws.append(["F10126106_A___DWG1", "Sheet one"])
    ws.append(["F10126107", "-"])
    ws.append(["not a stem row", "Junk"])
    wb.save(out)
    assert fx._prior_names(out) == {"F10126106_A___DWG1": "Sheet one"}


def test_prior_names_missing_or_corrupt_yields_empty(tmp_path):
    assert fx._prior_names(tmp_path / "nope.xlsx") == {}
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a workbook")
    assert fx._prior_names(bad) == {}


def test_prior_names_malformed_workbook_xml_yields_empty(tmp_path):
    # A zip-valid workbook with malformed inner XML raises
    # ElementTree.ParseError (a SyntaxError, not ValueError) on load: the
    # best-effort lookup yields {} instead of aborting the rebuild.
    import warnings
    import zipfile

    from openpyxl import Workbook
    out = tmp_path / "malformed-xml.xlsx"
    wb = Workbook()
    wb.active.title = "Files"
    wb.active.append(["Drawing (stem)", "Name"])
    wb.save(out)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with zipfile.ZipFile(out, "a") as zf:
            zf.writestr("xl/workbook.xml", b"<not xml")
    assert fx._prior_names(out) == {}


def test_prior_names_workbook_without_files_sheet_returns_empty(tmp_path):
    # A valid workbook from another tool has no Files sheet: the best-effort
    # lookup yields {} instead of aborting the rebuild with KeyError.
    from openpyxl import Workbook
    out = tmp_path / "foreign.xlsx"
    wb = Workbook()
    wb.active.title = "Summary"
    wb.active.append(["a", "b"])
    wb.save(out)
    assert fx._prior_names(out) == {}


def test_prior_names_malformed_files_rows_yield_empty(tmp_path, monkeypatch):
    # The Files sheet parses lazily: a read error during iter_rows (malformed
    # sheet XML) must yield {} instead of aborting the workbook rebuild.
    from openpyxl import Workbook
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    out = tmp_path / "malformed.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Files"
    ws.append(["Drawing (stem)", "Name"])
    ws.append(["F10126106", "Base name"])
    wb.save(out)

    def boom(self, *a, **k):
        raise ValueError("malformed sheet XML")

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", boom)
    assert fx._prior_names(out) == {}


def test_formula_injection_neutralized(tmp_path):
    # PDF-derived text starting with "=" must be stored as text, never as a
    # formula Excel evaluates on open (table rows + dashboard boxes).
    from openpyxl import Workbook, load_workbook

    out = tmp_path / "out"
    out.mkdir()
    path = tmp_path / "organize_fermi_report.xlsx"
    evil = '=HYPERLINK("http://evil.example","click")'
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 1, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [("F10126106", "F10126109")],
        "chk": [],
        "orphans": [],
        "roots": [],
        "mismatches": [],
        "report_txt": "",
        "names": {"F10126109": evil},
    }

    fx.build_workbook(path, ctx)
    cell = load_workbook(path)["Missing"]["B2"]
    assert cell.value == evil
    assert cell.data_type == "s"

    box = tmp_path / "box.xlsx"
    w2 = Workbook()
    fx._box(w2.active, 1, 1, 3, "=1+1")
    w2.save(box)
    cell2 = load_workbook(box).active["A1"]
    assert cell2.value == "=1+1"
    assert cell2.data_type == "s"


def test_dash_two_tier_banner_print_fit_and_scanned_info_color(tmp_path):
    from openpyxl import load_workbook

    def build(name, ctx):
        path = tmp_path / name
        fx.build_workbook(path, ctx)
        return load_workbook(path)

    base = {
        "output": str(tmp_path / "out"),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 0, "roots": 0, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [], "chk": [], "orphans": [], "roots": [], "mismatches": [],
        "report_txt": "", "names": {},
    }

    # Action tier only: Missing + CHK -> red "Action needed" banner.
    wb = build("act.xlsx", {**base, "missing": [("F10126106", "F10126109")],
                            "chk": ["F10126107__CHK"]})
    dash = wb["Dashboard"]
    assert dash.page_setup.fitToHeight == 1
    vals = [str(c.value) for row in dash.iter_rows() for c in row
            if c.value is not None]
    banner = next(v for v in vals if "Action needed" in v)
    assert "Worth a look" not in banner
    assert "Action" in vals          # per-row status pills
    assert "Info" in vals            # legend documents the blue swatch

    # Review tier only: one scanned PDF -> amber "Worth a look" + blue tab.
    wb2 = build("rev.xlsx", {**base, "scanned": [("F10126108", [1])]})
    vals2 = [str(c.value) for row in wb2["Dashboard"].iter_rows() for c in row
             if c.value is not None]
    assert any("Worth a look" in v and "Action needed" not in v for v in vals2)
    tab = wb2["Scanned"].sheet_properties.tabColor
    assert tab is not None and tab.rgb[-6:].upper() == "2563EB"


def test_dashboard_placement_refused_row(tmp_path):
    from openpyxl import load_workbook

    ctx = {
        "output": str(tmp_path / "out"),
        "run_mode": "DRY-RUN",
        "run_time": "2026-10-08T12:00:00",
        "counters": {"scanned": 4, "roots": 2, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [], "chk": [], "orphans": [], "roots": [], "mismatches": [],
        "report_txt": "", "names": {}, "placement_refused": True,
    }
    path = tmp_path / "refused.xlsx"
    fx.build_workbook(path, ctx)

    dash = load_workbook(path)["Dashboard"]
    vals = [str(c.value) for row in dash.iter_rows() for c in row
            if c.value is not None]
    assert "Placement refused" in vals
    # Refusal is action-tier: the banner must not read "All clear".
    assert any("Action needed" in v for v in vals)
    assert not any("All clear" in v for v in vals)


def test_refresh_from_run_forwards_placement_refused(tmp_path):
    # The CLI path reaches the Dashboard through refresh_from_run: a refused
    # run must light the action row (previously the flag was dropped here,
    # so the row always read 0 after real runs).
    from openpyxl import load_workbook

    out = tmp_path / "out"
    xlsx = fx.refresh_from_run(
        output=str(out), run_mode="DRY-RUN", run_time="2026-10-08T12:00:00",
        counters={"scanned": 1, "roots": 1, "copies": 0, "cycles": 0,
                  "warnings": 0},
        missing=[], chk=[], orphans=[], report_txt=None,
        placement_refused=True)

    vals = [str(c.value) for row in load_workbook(xlsx)["Dashboard"].iter_rows()
            for c in row if c.value is not None]
    assert "Placement refused" in vals
    assert any("Action needed" in v for v in vals)


def test_no_formula_cells_anywhere_including_used_on_bugs(tmp_path):
    # Fuzz-style: a "=" payload in the USED ON BUGS section (written outside
    # _table) must never be persisted as an Excel formula, and no other sheet
    # may produce one either.
    from openpyxl import load_workbook

    out = tmp_path / "out"
    out.mkdir()
    path = tmp_path / "organize_fermi_report.xlsx"
    evil = "=HYPERLINK(\"http://evil.example\",\"hi\")"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 1, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [], "chk": [], "orphans": [], "roots": [], "mismatches": [],
        "used_on_bugs": [(evil, evil)],
        "report_txt": "",
        "names": {evil: evil},
    }

    fx.build_workbook(path, ctx)
    wb = load_workbook(path)
    formulas = [(ws.title, c.coordinate)
                for ws in wb.worksheets for row in ws.iter_rows()
                for c in row if c.data_type == "f"]
    assert formulas == []


def test_formula_fuzz_every_input_field_produces_zero_formulas(tmp_path):
    # Full-field fuzz: a "=" payload in every ctx input (including the
    # Dashboard provenance via _box, KPI card counter values, and the Run
    # History row) must produce zero formula cells anywhere in the workbook.
    # Pins every dynamic write path so a future edit cannot drop a guard.
    from openpyxl import load_workbook

    evil = "=HYPERLINK(\"http://evil.example\",\"x\")"
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": evil,              # Dashboard provenance (via _box)
        "run_mode": evil,            # Dashboard subtitle + Run History row
        "run_time": evil,
        "counters": {"scanned": evil, "roots": evil, "copies": evil,
                     "cycles": evil, "warnings": evil},   # KPI card values (via _box)
        "missing": [(evil, evil), ("F10126106", "F10126109")],
        "chk": [evil],
        "orphans": [evil],
        "roots": [(evil, [evil])],
        "mismatches": [(evil, evil, [evil])],
        "used_on_bugs": [(evil, evil)],
        "titleblock_mismatches": [(evil, "revision", evil, evil)],
        "scanned": [(evil, [1])],
        "watermarks": [(evil, evil)],
        "report_txt": evil,          # Run History report column
        "names": {evil: evil, "F10126109": evil},
    }

    fx.build_workbook(path, ctx)
    wb = load_workbook(path)
    formulas = [(ws.title, c.coordinate)
                for ws in wb.worksheets for row in ws.iter_rows()
                for c in row if c.data_type == "f"]
    assert formulas == []
    # Guards are only pinned if the payload actually flows into these sheets:
    # an omitted dynamic field would otherwise pass vacuously above.
    for sheet in ("Dashboard", "Run History"):
        hits = [c for row in wb[sheet].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
        assert hits, f"fuzz payload must reach {sheet} cells"
        assert all(c.data_type == "s" for c in hits), sheet


def test_formula_injection_extended_prefixes(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "out"
    out.mkdir()
    path = tmp_path / "organize_fermi_report.xlsx"
    prefixes = ["=", "@", "+", "-", "\t", "\r"]
    missing_items = [(f"STEM_{i}", f"{p}evil_payload") for i, p in enumerate(prefixes)]
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 0, "copies": 0, "cycles": 0, "warnings": 0},
        "missing": missing_items,
        "chk": [], "orphans": [], "roots": [], "mismatches": [],
        "report_txt": "", "names": {},
    }

    fx.build_workbook(path, ctx)
    wb = load_workbook(path)
    ws = wb["Missing"]
    for row_idx, p in enumerate(prefixes, start=2):
        cell = ws.cell(row=row_idx, column=1)
        assert cell.data_type == "s"
        if p == "\r":
            assert cell.value in ("\revil_payload", "\nevil_payload")
        else:
            assert cell.value == f"{p}evil_payload"


def test_link_target_length_guard():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    c1 = ws.cell(row=1, column=1, value="normal")
    fx._link(c1, "http://example.com/short")
    assert c1.hyperlink is not None
    assert c1.hyperlink.target == "http://example.com/short"

    c2 = ws.cell(row=2, column=1, value="long")
    long_target = "http://example.com/" + ("a" * 250)
    fx._link(c2, long_target)
    assert c2.hyperlink is None

    c3 = ws.cell(row=3, column=1, value="none")
    fx._link(c3, None)
    assert c3.hyperlink is None


def test_mismatch_sheet_dynamic_layout_spacing(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "out"
    out.mkdir()

    # Case 1: Empty mismatches data -> table1_rows = 1 (empty state row) -> base = 5
    path_empty = tmp_path / "empty_mism.xlsx"
    ctx_empty = {
        "output": str(out), "run_mode": "DRY-RUN", "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 0, "roots": 0, "copies": 0, "cycles": 0, "warnings": 0},
        "missing": [], "chk": [], "orphans": [], "roots": [], "mismatches": [],
        "used_on_bugs": [("P1", "C1")], "report_txt": "", "names": {},
    }
    fx.build_workbook(path_empty, ctx_empty)
    ws_empty = load_workbook(path_empty)["USED ON check"]
    assert ws_empty.cell(row=5, column=1).value == "USED ON BUGS (1)"

    # Case 2: 2 mismatches -> table1_rows = 2 -> base = 6
    path_populated = tmp_path / "populated_mism.xlsx"
    ctx_pop = {
        "output": str(out), "run_mode": "DRY-RUN", "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 0, "roots": 0, "copies": 0, "cycles": 0, "warnings": 0},
        "missing": [], "chk": [], "orphans": [], "roots": [],
        "mismatches": [("P1", "C1", ["A1"]), ("P2", "C2", ["A2"])],
        "used_on_bugs": [("P1", "C1")], "report_txt": "", "names": {},
    }
    fx.build_workbook(path_populated, ctx_pop)
    ws_pop = load_workbook(path_populated)["USED ON check"]
    assert ws_pop.cell(row=6, column=1).value == "USED ON BUGS (1)"


def test_workbook_context_typeddict():
    ctx: fx.WorkbookContext = {
        "output": "/tmp/out",
        "run_mode": "EXECUTE",
        "run_time": "2026-10-15T10:00:00",
        "counters": {"scanned": 5},
        "missing": [],
        "chk": [],
        "orphans": [],
        "roots": [],
        "mismatches": [],
        "used_on_bugs": [],
        "titleblock_mismatches": [],
        "scanned": [],
        "watermarks": [],
        "report_txt": "report.txt",
        "names": {},
        "placement_refused": False,
    }
    assert ctx["output"] == "/tmp/out"
    assert ctx.get("run_mode") == "EXECUTE"


def test_second_run_renders_the_history_trend_chart(tmp_path):
    # The n>=2 block (trend note + LineChart) only runs on a relaunch against an
    # existing workbook; a defect there used to ship green because every test
    # built a first-run workbook.
    import fermi_report_xlsx as fx

    out = tmp_path / "out"
    out.mkdir()
    fx.refresh_from_run(out, "EXECUTE", "2026-10-26 10:00", {"scanned": 3, "copies": 2},
                        [], [], [], "report one")
    fx.refresh_from_run(out, "EXECUTE", "2026-10-26 11:00", {"scanned": 4, "copies": 3},
                        [], [], [], "report two")

    xlsx = out / "organize_fermi_report.xlsx"
    from openpyxl import load_workbook

    ws = load_workbook(xlsx)["Run History"]
    assert len(ws._charts) == 1
    assert ws.max_row >= 3


def test_diagnostic_sheets_link_the_file_they_name(tmp_path):
    # Every sheet except Missing names a PDF that exists in the tree, so its
    # stem must be clickable too - not only the Files / Orphans / Superseded
    # listings that carry a path.
    out = tmp_path / "out"
    (out / "F10126106 Assembly").mkdir(parents=True)
    (out / "F10126106 Assembly" / "F10126106.pdf").write_bytes(b"%PDF")
    (out / "_superseded").mkdir()
    (out / "_superseded" / "F10126107__CHK.pdf").write_bytes(b"%PDF")
    (out / "_orphans").mkdir()
    (out / "_orphans" / "F10126108.pdf").write_bytes(b"%PDF")
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 3, "roots": 1, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [("F10126109", "F10126106")],
        "chk": ["F10126107__CHK"],
        "orphans": ["F10126108"],
        "roots": [("F10126106", ["F10126100"])],
        "mismatches": [("F10126106", "F10126108", ["F10126100"])],
        "used_on_bugs": [("F10126106", "F10126108")],
        "titleblock_mismatches": [("F10126106", "revision", "-", "B")],
        "scanned": [("F10126106", [1])],
        "watermarks": [("F10126108", "watermark text: DRAFT")],
        "report_txt": "",
        "names": {},
    }

    fx.build_workbook(path, ctx)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    for sheet in ("CHK", "Scanned", "Watermark", "Roots", "Title block check",
                  "USED ON check"):
        assert wb[sheet]["A2"].hyperlink is not None, f"{sheet} did not link its stem"
    # The target is the file the row names, not just some link.
    assert wb["CHK"]["A2"].hyperlink.target == "_superseded/F10126107__CHK.pdf"
    # USED ON check names two drawings; both columns are clickable.
    assert wb["USED ON check"]["C2"].hyperlink.target == "_orphans/F10126108.pdf"
    # Missing exists precisely because there is no file to open.
    assert wb["Missing"]["A2"].hyperlink is None


def test_diagnostic_links_prefer_the_active_copy_over_the_archive(tmp_path):
    # A superseded twin shares the canonical stem, and _superseded/ sorts
    # before the drawing folder: the link must still open the live revision.
    out = tmp_path / "out"
    (out / "F10126106 Assembly").mkdir(parents=True)
    (out / "F10126106 Assembly" / "F10126106___DWG1.pdf").write_bytes(b"%PDF")
    (out / "_superseded").mkdir()
    (out / "_superseded" / "F10126106---DWG1.pdf").write_bytes(b"%PDF")
    path = tmp_path / "organize_fermi_report.xlsx"
    ctx = {
        "output": str(out),
        "run_mode": "DRY-RUN",
        "run_time": "2026-09-30T12:00:00",
        "counters": {"scanned": 1, "roots": 1, "copies": 0,
                     "cycles": 0, "warnings": 0},
        "missing": [],
        "chk": [],
        "orphans": [],
        "roots": [("F10126106___DWG1", [])],
        "mismatches": [],
        "report_txt": "",
        "names": {},
    }

    fx.build_workbook(path, ctx)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert (wb["Roots"]["A2"].hyperlink.target
            == "F10126106 Assembly/F10126106___DWG1.pdf")

