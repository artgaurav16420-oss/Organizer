"""Live Excel report for the Fermi PDF organizer.

Rebuilt after every run from the run's counters + a live disk scan of the output
folder, so the workbook always shows the current status of the tree.
Sheets: Dashboard, Missing, CHK, Scanned, Watermark, Orphans, Superseded,
Roots, USED ON check, Title block check, Files (live tree), Run History.
Run History survives rebuilds via a sidecar JSON.

Layout contract (pinned by tests): every list sheet has its header in row 1 and
data from row 2; Dashboard card 1 is label B5 / value B6.

Values are written by the organizer at run time (a snapshot, not formulas);
Files / Orphans / Superseded counts come from a live disk scan at that moment.

openpyxl is optional; when it is missing the workbook is skipped gracefully.
"""
import json
import os
import shutil
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint
    from openpyxl.formatting.rule import DataBarRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
    _OPENPYXL_OK = True
except ImportError:
    _OPENPYXL_OK = False

from fermi_organizer.config import canonical_stem  # noqa: E402
from fermi_organizer.fsops import scan_output_tree  # noqa: E402

# ---------------------------------------------------------------- design tokens
FONT = "Arial"
DARK = "1F3864"        # header navy
NAVY = "2E5AAC"        # links / accents
INK = "1F2937"         # body text
MUTED = "6B7280"       # secondary text
LINE = "E5E7EB"        # hairlines
STRIPE = "F8FAFC"      # zebra rows
PANEL = "F1F5F9"       # note panels / card-less areas
WHITE = "FFFFFF"

# Severity palettes: tint fill, strong text, accent (tab / card edge / chart).
PAL = {
    "red":    {"fill": "FDE2E1", "text": "9B1C1C", "accent": "DC2626"},   # action required
    "amber":  {"fill": "FEF3C7", "text": "92400E", "accent": "D97706"},   # review
    "orange": {"fill": "FFE8D5", "text": "9A3412", "accent": "EA580C"},   # parked
    "grey":   {"fill": "E5E7EB", "text": "374151", "accent": "6B7280"},   # archived
    "green":  {"fill": "DCFCE7", "text": "166534", "accent": "16A34A"},   # ok / active
    "blue":   {"fill": "DBEAFE", "text": "1E40AF", "accent": "2563EB"},   # info
}

# Back-compat color aliases (fill, text) for any external reader of the module.
RED_F, RED_D = PAL["red"]["fill"], PAL["red"]["text"]
AMB_F, AMB_D = PAL["amber"]["fill"], PAL["amber"]["text"]
ORG_F, ORG_D = PAL["orange"]["fill"], PAL["orange"]["text"]
GREY_F, GREY_D = PAL["grey"]["fill"], PAL["grey"]["text"]
GRN_F, GRN_D = PAL["green"]["fill"], PAL["green"]["text"]

HAIR = Side(style="thin", color=LINE)
ROW_BORDER = Border(bottom=HAIR)
PILL_EDGE = Side(style="medium", color=WHITE)

# Run-history cap: single source for both the sidecar save and the sheet render
# (the two must agree or the saved file and the displayed table drift).
HISTORY_MAX_RUNS = 200
MAX_COL_WIDTH = 64
ROW_H = 21


# -------------------------------------------------------------------- helpers
def _font(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color,
                underline=underline)


def _solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def _page_setup(ws):
    """Landscape, fit to one page wide, header row repeats when printed."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def _hdr(ws, headers, widths, tab=None, aligns=None):
    """Header row 1 + freeze panes + widths + no gridlines + tab color."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    if tab:
        ws.sheet_properties.tabColor = PAL[tab]["accent"] if tab in PAL else tab
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = _font(10, True, WHITE)
        c.fill = _solid(DARK)
        horiz = (aligns[i - 1] if aligns else "left")
        c.alignment = Alignment(horizontal=horiz, vertical="center",
                                wrap_text=True, indent=1 if horiz == "left" else 0)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    _page_setup(ws)
    ws.print_title_rows = "1:1"


def _pill(cell, palette):
    """Status pill: tinted fill, bold strong text, white inset edge."""
    p = PAL[palette]
    cell.fill = _solid(p["fill"])
    cell.font = _font(9, True, p["text"])
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(left=PILL_EDGE, right=PILL_EDGE, top=PILL_EDGE, bottom=PILL_EDGE)


def _link(cell, target, bold=False):
    """Hyperlink styled as a link; failures leave the plain value."""
    try:
        cell.hyperlink = target
        cell.font = _font(10, bold, NAVY, underline="single")
    except Exception:  # pragma: no cover - openpyxl rejects only malformed targets
        pass


def _file_uri(path):
    try:
        return Path(path).absolute().as_uri()
    except (ValueError, OSError):  # pragma: no cover - relative/odd paths
        return None


def _rel(path, output):
    try:
        return str(Path(path).relative_to(Path(output)))
    except (ValueError, TypeError):
        return Path(path).name


def _autofit(ws, ncols, nrows, minimums, cap=MAX_COL_WIDTH):
    """Widen columns to their content (rows 2..nrows+1), capped; never shrink."""
    for ci in range(1, ncols + 1):
        longest = 0
        for r in range(2, nrows + 2):
            v = ws.cell(row=r, column=ci).value
            if v is not None:
                longest = max(longest, len(str(v)))
        want = min(cap, longest + 4)
        cur = minimums[ci - 1] if ci - 1 < len(minimums) else 10
        ws.column_dimensions[get_column_letter(ci)].width = max(cur, want)


def _side_note(ws, ncols, title, text, palette="blue"):
    """About-this-sheet panel right of the table, plus a way back."""
    col = ncols + 2
    L = get_column_letter(col)
    ws.column_dimensions[get_column_letter(ncols + 1)].width = 3
    ws.column_dimensions[L].width = 46
    p = PAL[palette]
    h = ws.cell(row=1, column=col, value=title)
    h.font = _font(10, True, p["text"])
    h.fill = _solid(p["fill"])
    h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(start_row=2, start_column=col, end_row=7, end_column=col)
    body = ws.cell(row=2, column=col, value=text)
    body.font = _font(9, False, INK)
    body.fill = _solid(PANEL)
    body.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    back = ws.cell(row=8, column=col, value="\u2190 Back to Dashboard")
    _link(back, "#'Dashboard'!A1", bold=True)
    back.alignment = Alignment(vertical="center", indent=1)


def _empty_state(ws, ncols, message):
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=message)
    c.font = _font(10, True, PAL["green"]["text"])
    for ci in range(1, ncols + 1):
        ws.cell(row=2, column=ci).fill = _solid(PAL["green"]["fill"])
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 28


def _table(ws, headers, widths, rows, *, tab, aligns=None, pills=None,
           link_first=None, empty="Nothing to show", note=None, first_bold=True,
           max_width=MAX_COL_WIDTH):
    """Generic list sheet: header row 1, zebra rows from row 2, optional pills.

    rows: list of value lists.  pills: {col_index: palette | callable(row)->palette}
    link_first: optional callable(row_index, row)->target for column 1.
    """
    aligns = aligns or ["left"] * len(headers)
    _hdr(ws, headers, widths, tab=tab, aligns=aligns)
    ncols = len(headers)
    if not rows:
        _empty_state(ws, ncols, f"\u2713  {empty}")
    for ri, vals in enumerate(rows, start=2):
        ws.row_dimensions[ri].height = ROW_H
        stripe = STRIPE if ri % 2 == 1 else WHITE
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.fill = _solid(stripe)
            c.border = ROW_BORDER
            c.font = _font(10, first_bold and ci == 1, INK)
            horiz = aligns[ci - 1]
            # No wrapping: rows stay one line tall; the last column may overflow
            # into the empty space to its right.
            c.alignment = Alignment(horizontal=horiz, vertical="center",
                                    indent=1 if horiz == "left" else 0)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                c.number_format = "#,##0"
        if pills:
            for ci, pal in pills.items():
                key = pal(vals) if callable(pal) else pal
                _pill(ws.cell(row=ri, column=ci), key)
        if link_first:
            target = link_first(ri - 2, vals)
            if target:
                _link(ws.cell(row=ri, column=1), target, bold=True)
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{len(rows) + 1}"
        _autofit(ws, ncols, len(rows), widths, max_width)
    if note:
        _side_note(ws, ncols, note[0], note[1], note[2] if len(note) > 2 else "blue")
    return ncols


def _scan_tree(output):
    """Live disk scan of the output folder -> {'tree','sup','orph'} path lists."""
    return scan_output_tree(output)


def _hist_path(xlsx_path):
    return Path(str(xlsx_path) + ".history.json")


def _load_history(xlsx_path, log=None):
    hist = _hist_path(xlsx_path)
    if not hist.exists():
        # A sidecar that was never written is normal, not a warning.
        return []
    try:
        data = json.loads(hist.read_text(encoding="utf-8"))
        if not isinstance(data, list):
            raise ValueError("run history is not a list")
        return data
    except Exception as e:
        (log or print)(f"WARNING: run history not loaded ({hist}): {e}")
        if hist.exists():
            try:
                shutil.copy2(hist, str(hist) + ".corrupt")
            except OSError:
                pass
        return []


def _save_history(xlsx_path, rows, log=None):
    hist = _hist_path(xlsx_path)
    tmp = hist.with_name(hist.name + ".tmp")
    try:
        tmp.write_text(json.dumps(rows[-HISTORY_MAX_RUNS:]), encoding="utf-8")
        os.replace(tmp, hist)
    except OSError as e:
        (log or print)(f"WARNING: run history not saved: {e}")


def _resolve_names(names_raw, live):
    """names_raw + base-stem fallback, for every tree/orphan stem and every key."""
    resolved = {}
    live_stems = [canonical_stem(q.stem) for q in live["tree"] + live["orph"]]
    for stem in set(list(names_raw) + [s for s in live_stems if s]):
        v = names_raw.get(stem)
        if not v:
            v = names_raw.get(stem.split("_")[0], "")
        if v:
            resolved[stem] = v
    return resolved



def build_workbook(path, ctx, log=None):
    """Build/refresh the live workbook at `path`; returns it (None without openpyxl).

    Context keys read (15): output, run_mode, run_time,
    counters{scanned,roots,copies,cycles,warnings}, missing[(stem,val)...],
    chk[stem...], orphans[stem...], roots[(stem,[used_on])...],
    mismatches[(parent,child,actual)...], used_on_bugs[(parent,child)...],
    titleblock_mismatches[(stem,field,filename,tb)...],
    scanned[(stem,[pages])...], watermarks[(stem,evidence)...],
    report_txt, names{stem/val: NAME}.
    cli.py maps the run-context key used_on_mismatches -> mismatches and
    used_on_bugs -> used_on_bugs when it calls refresh_from_run.
    """
    if not _OPENPYXL_OK:
        return None
    path = Path(path)
    wb = Workbook()
    wb.properties.title = "Fermi PDF Organizer - live report"
    wb.properties.creator = "Fermi PDF Organizer"
    dash = wb.active
    dash.title = "Dashboard"
    missing = wb.create_sheet("Missing")
    chk = wb.create_sheet("CHK")
    scanned_ws = wb.create_sheet("Scanned")
    watermark_ws = wb.create_sheet("Watermark")
    orph = wb.create_sheet("Orphans")
    superseded = wb.create_sheet("Superseded")
    roots_ws = wb.create_sheet("Roots")
    mism_ws = wb.create_sheet("USED ON check")
    tb_ws = wb.create_sheet("Title block check")
    files = wb.create_sheet("Files")
    hist = wb.create_sheet("Run History")

    counters = ctx.get("counters", {})
    live = _scan_tree(ctx["output"])
    missing_pairs = sorted(set(ctx.get("missing", [])))
    names = _resolve_names(ctx.get("names", {}) or {}, live)

    _dash_sheet(dash, ctx, counters, live, missing_pairs)
    _missing_sheet(missing, missing_pairs, names)
    _chk_sheet(chk, ctx, names)
    _scanned_sheet(scanned_ws, ctx, names)
    _watermark_sheet(watermark_ws, ctx, names)
    _orphans_sheet(orph, ctx, names, live)
    _superseded_sheet(superseded, live)
    _roots_sheet(roots_ws, ctx, names)
    _mismatch_sheet(mism_ws, ctx, names)
    _titleblock_sheet(tb_ws, ctx, names)
    _files_sheet(files, ctx, names, live)
    _history_sheet(hist, ctx, counters, missing_pairs, path, log)

    wb.active = 0
    wb.save(path)
    return path


# ------------------------------------------------------------------ Dashboard
_CARD_SPANS = ((2, 4), (6, 8), (10, 12), (14, 16))     # B:D  F:H  J:L  N:P


def _box(ws, row, c1, c2, value=None, font=None, fill=None, align=None,
         border=None):
    """Merge c1..c2 on one row; style every cell so the block renders solid."""
    for cc in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=cc)
        if fill:
            cell.fill = fill
        if border:
            cell.border = border
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    a = ws.cell(row=row, column=c1)
    a.value = value
    if font:
        a.font = font
    if align:
        a.alignment = align
    return a


def _card(ws, row, span, label, value, palette, caption, target):
    c1, c2 = span
    p = PAL[palette]
    edge = Side(style="thick", color=p["accent"])
    fill = _solid(p["fill"])
    left = Border(left=edge)
    _box(ws, row, c1, c2, label, _font(8, True, MUTED), fill,
         Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1), left)
    _box(ws, row + 1, c1, c2, value, _font(24, True, p["text"]), fill,
         Alignment(horizontal="left", vertical="center", indent=1), left)
    cap = _box(ws, row + 2, c1, c2, caption, _font(9, False, MUTED), fill,
               Alignment(horizontal="left", vertical="center", indent=1), left)
    if target:
        _link(cap, target)
    ws.row_dimensions[row].height = 30
    ws.row_dimensions[row + 1].height = 40
    ws.row_dimensions[row + 2].height = 20


def _dash_sheet(dash, ctx, counters, live, missing_pairs):
    dash.sheet_view.showGridLines = False
    dash.sheet_properties.tabColor = DARK
    dash.sheet_view.zoomScale = 100
    _page_setup(dash)
    widths = {"A": 2, "E": 2, "I": 2, "M": 2, "Q": 2}
    for col in "BCDFGHJKLNOP":
        widths[col] = 9.5
    for col, w in widths.items():
        dash.column_dimensions[col].width = w

    # ---- title block
    dash.row_dimensions[1].height = 10
    dash.row_dimensions[2].height = 34
    t = _box(dash, 2, 2, 16, "Fermi PDF Organizer", _font(22, True, DARK),
             None, Alignment(horizontal="left", vertical="center"))
    t.value = "Fermi PDF Organizer"
    mode = str(ctx.get("run_mode", ""))
    sub = f"Live status  \u00b7  last run {ctx.get('run_time', '')}  \u00b7  {mode}"
    _box(dash, 3, 2, 16, sub, _font(10, False, MUTED), None,
         Alignment(horizontal="left", vertical="center"))
    dash.row_dimensions[3].height = 20
    dash.row_dimensions[4].height = 10

    sup_n, orph_n = len(live["sup"]), len(live["orph"])
    chk_n = len(set(ctx.get("chk", [])))
    warn = counters.get("warnings", "-")
    cyc = counters.get("cycles", 0)
    cards = [
        ("PDF FILES IN TREE (LIVE)", len(live["tree"]), "green", "Open Files \u2192", "Files"),
        ("ROOTS (LAST RUN)", counters.get("roots", "-"), "green", "Open Roots \u2192", "Roots"),
        ("MISSING - DOWNLOAD FROM TEAMCENTER", len(missing_pairs),
         "red" if missing_pairs else "green", "Open Missing \u2192", "Missing"),
        ("CHK UNAPPROVED", chk_n, "amber" if chk_n else "green", "Open CHK \u2192", "CHK"),
        ("ORPHANS (PARKED ON FULL RUN)", orph_n, "orange" if orph_n else "green",
         "Open Orphans \u2192", "Orphans"),
        ("SUPERSEDED REVISIONS (LIVE)", sup_n, "grey", "Open Superseded \u2192", "Superseded"),
        ("EXTRACTION WARNINGS (LAST RUN)", warn,
         "amber" if isinstance(warn, int) and warn else "green", "Details in report .txt", None),
        ("CYCLES BROKEN (LAST RUN)", cyc, "green" if not cyc else "amber",
         "Cycles are cut automatically", None),
    ]
    for i, (label, val, pal, cap, sheet) in enumerate(cards):
        row = 5 + (i // 4) * 4
        _card(dash, row, _CARD_SPANS[i % 4], label, val, pal, cap,
              f"#'{sheet}'!A1" if sheet else None)
    dash.row_dimensions[8].height = 12
    dash.row_dimensions[12].height = 14

    # ---- needs attention
    items = [
        ("Missing parts", len(missing_pairs), "red", "Missing"),
        ("CHK unapproved", chk_n, "amber", "CHK"),
        ("Orphans (parent not seen)", orph_n, "orange", "Orphans"),
        ("Watermarked PDFs", len(ctx.get("watermarks", [])), "amber", "Watermark"),
        ("Scanned PDFs (OCR)", len(ctx.get("scanned", [])), "orange", "Scanned"),
        ("USED ON check", len(ctx.get("mismatches", [])) + len(ctx.get("used_on_bugs", [])),
         "amber", "USED ON check"),
        ("Title block check", len(ctx.get("titleblock_mismatches", [])), "amber",
         "Title block check"),
    ]
    todo = sum(1 for _, n, _, _ in items if n)
    r = 13
    _box(dash, r, 2, 16, "NEEDS ATTENTION", _font(11, True, DARK), None,
         Alignment(horizontal="left", vertical="center"),
         Border(bottom=Side(style="medium", color=DARK)))
    dash.row_dimensions[r].height = 26
    ok = todo == 0
    banner = ("\u2713  All clear - nothing needs review" if ok else
              f"\u26a0  {todo} of {len(items)} categories need review")
    _box(dash, r + 1, 2, 16, banner,
         _font(11, True, PAL["green" if ok else "amber"]["text"]),
         _solid(PAL["green" if ok else "amber"]["fill"]),
         Alignment(horizontal="left", vertical="center", indent=1))
    dash.row_dimensions[r + 1].height = 30
    dash.row_dimensions[r + 2].height = 8

    h = r + 3
    for (c1, c2), text, horiz in (((2, 4), "Category", "left"), ((6, 8), "Count", "center"),
                                  ((10, 12), "Status", "center"), ((14, 16), "Go to", "left")):
        _box(dash, h, c1, c2, text, _font(9, True, WHITE), _solid(DARK),
             Alignment(horizontal=horiz, vertical="center", indent=1 if horiz == "left" else 0))
    dash.row_dimensions[h].height = 24
    first = h + 1
    for k, (label, n, pal, sheet) in enumerate(items):
        rr = first + k
        stripe = _solid(STRIPE if k % 2 else WHITE)
        dash.row_dimensions[rr].height = 22
        _box(dash, rr, 2, 4, label, _font(10, True, INK), stripe,
             Alignment(horizontal="left", vertical="center", indent=1), ROW_BORDER)
        cnt = _box(dash, rr, 6, 8, n, _font(11, True, INK), stripe,
                   Alignment(horizontal="center", vertical="center"), ROW_BORDER)
        cnt.number_format = "#,##0"
        st = _box(dash, rr, 10, 12, "Review" if n else "OK", None, None, None, None)
        for cc in range(10, 13):
            dash.cell(row=rr, column=cc).fill = _solid(PAL[pal if n else "green"]["fill"])
        _pill(st, pal if n else "green")
        go = _box(dash, rr, 14, 16, f"Open {sheet} \u2192", None, stripe,
                  Alignment(horizontal="left", vertical="center", indent=1), ROW_BORDER)
        _link(go, f"#'{sheet}'!A1")
    last = first + len(items) - 1

    # ---- chart (native, editable in Excel)
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = "Items needing review"
    ch.legend = None
    ch.height, ch.width = 7.5, 24
    ch.y_axis.majorGridlines = None
    ch.x_axis.delete = False
    ch.y_axis.delete = True            # data labels carry the values
    ch.x_axis.scaling.orientation = "maxMin"
    ch.add_data(Reference(dash, min_col=6, min_row=first, max_row=last), titles_from_data=False)
    ch.set_categories(Reference(dash, min_col=2, min_row=first, max_row=last))
    series = ch.series[0]
    series.graphicalProperties.solidFill = PAL["blue"]["accent"]
    for idx, (_, n, pal, _) in enumerate(items):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = PAL[pal if n else "green"]["accent"]
        series.dPt.append(pt)
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.showSerName = ch.dataLabels.showCatName = ch.dataLabels.showLegendKey = False
    dash.add_chart(ch, f"B{last + 3}")

    # ---- footer: legend + provenance
    f = last + 3 + 17
    _box(dash, f, 2, 16, "COLOR GUIDE", _font(9, True, MUTED), None,
         Alignment(horizontal="left", vertical="center"))
    legend = (("red", "Action required"), ("amber", "Review"), ("orange", "Parked"),
              ("grey", "Archived"), ("green", "OK / active"))
    spans = ((2, 3), (4, 6), (7, 9), (10, 12), (13, 15))
    for (c1, c2), (pal, text) in zip(spans, legend):
        for cc in range(c1, c2 + 1):
            dash.cell(row=f + 1, column=cc).fill = _solid(PAL[pal]["fill"])
        _pill(_box(dash, f + 1, c1, c2, text, None, None, None, None), pal)
    dash.row_dimensions[f + 1].height = 22
    out = ctx.get("output", "")
    rep = ctx.get("report_txt", "")
    note = f"Output: {out}" + (f"   \u00b7   Report: {rep}" if rep else "")
    _box(dash, f + 3, 2, 16, note, _font(9, False, MUTED, italic=True), None,
         Alignment(horizontal="left", vertical="center", wrap_text=True))
    _box(dash, f + 4, 2, 16,
         "Rebuilt after every run. Counts for Files / Orphans / Superseded come from a "
         "live disk scan at that moment; all other figures are from the last run.",
         _font(9, False, MUTED, italic=True), None,
         Alignment(horizontal="left", vertical="center", wrap_text=True))
    dash.row_dimensions[f + 4].height = 28


# ----------------------------------------------------------------- list sheets
def _missing_sheet(missing, missing_pairs, names):
    rows = [[val, names.get(val, "") or "-", stem, "Download approved DWG from Teamcenter"]
            for stem, val in missing_pairs]
    _table(missing, ["Missing FERMI#", "Name", "Referenced by drawing", "Action"],
           (22, 38, 30, 42), rows, tab="red", pills={4: "red"},
           empty="No missing parts - every referenced drawing is in the folder",
           note=("About this sheet",
                 "Part numbers referenced by a parent's BOM that have no PDF in the input "
                 "folder. Download the approved DWG from Teamcenter and rerun "
                 "(incremental is enough).", "red"))


def _chk_sheet(chk, ctx, names):
    rows = [[s, names.get(s, "") or "-", "Unapproved",
             "Replace with approved DWG from Teamcenter"]
            for s in sorted(set(ctx.get("chk", [])))]
    _table(chk, ["Stem (current file)", "Name", "Status", "Action"],
           (30, 38, 16, 48), rows, tab="amber", pills={3: "amber"},
           aligns=["left", "left", "center", "left"],
           empty="No unapproved CHK drawings",
           note=("About this sheet",
                 "CHK drawings are unapproved check prints. When the approved DWG arrives "
                 "in an input folder it supersedes the CHK automatically.", "amber"))


def _scanned_sheet(scanned_ws, ctx, names):
    data = ctx.get("scanned", [])  # [(stem, [pages]), ...]
    rows = []
    for stem, pages in sorted(data):
        pages_txt = (", ".join(str(p) for p in pages)
                     if isinstance(pages, (list, tuple)) else str(pages))
        rows.append([stem, names.get(stem, "") or "-", pages_txt])
    _table(scanned_ws, ["Stem (scanned PDF)", "Name", "OCR page(s)"], (30, 44, 18), rows,
           tab="orange", aligns=["left", "left", "center"], pills={3: "orange"},
           empty="No scanned (image-only) PDFs",
           note=("About this sheet",
                 "Image-only (scanned) PDFs whose text was read via OCR. OCR is "
                 "best-effort: verify the BOM of these drawings by eye.", "orange"))


def _watermark_sheet(watermark_ws, ctx, names):
    data = ctx.get("watermarks", [])  # [(stem, evidence), ...]
    rows = [[stem, names.get(stem, "") or "-", evidence or "-"]
            for stem, evidence in sorted(data)]
    _table(watermark_ws, ["Stem (watermarked PDF)", "Name", "Evidence"], (30, 40, 52), rows,
           tab="amber", empty="No watermarked PDFs detected",
           note=("About this sheet",
                 "Best-effort detection: watermark layer, transparent/diagonal/light "
                 "large text, or a watermark keyword.", "amber"))


def _orphans_sheet(orph, ctx, names, live=None):
    paths = {}
    for p in (live or {}).get("orph", []):
        paths[canonical_stem(p.stem) or p.stem.upper()] = p
    stems = sorted(ctx.get("orphans", []))
    rows = [[s, names.get(s, "") or "-", "Parked",
             "Parent drawing not seen so far - adopted into tree once it arrives"]
            for s in stems]
    _table(orph, ["Stem (current file)", "Name", "Status", "Note"], (30, 38, 14, 60), rows,
           tab="orange", pills={3: "orange"}, aligns=["left", "left", "center", "left"],
           link_first=lambda i, row: _file_uri(paths[row[0]]) if row[0] in paths else None,
           empty="No orphans - every part has a parent",
           note=("About this sheet",
                 "Orphans are stored under _orphans/ so a future incremental run can "
                 "adopt them when their parent appears. Click a stem to open its PDF.",
                 "orange"))


def _superseded_sheet(superseded, live):
    ordered = sorted(live["sup"], key=lambda q: q.name.lower())
    rows = []
    for p in ordered:
        stem = canonical_stem(p.stem) or p.stem.upper()
        rows.append([p.name, stem.split("_")[0], "Archived",
                     "older revision / unapproved CHK / superseded in place"])
    _table(superseded, ["File in _superseded/", "Base", "Status", "Note"], (44, 18, 14, 52),
           rows, tab="grey", pills={3: "grey"}, aligns=["left", "left", "center", "left"],
           link_first=lambda i, row: _file_uri(ordered[i]),
           empty="Nothing superseded yet",
           note=("About this sheet",
                 "Older revisions, unapproved CHK prints and in-place superseded "
                 "copies are archived here, never deleted. Click a name to open.", "grey"))


def _roots_sheet(roots_ws, ctx, names):
    data = ctx.get("roots", [])  # [(stem, [used_on...]), ...]
    rows = [[stem, names.get(stem, ""), ", ".join(sorted(set(used))) or "-"]
            for stem, used in data]
    _table(roots_ws, ["Root drawing", "Name", "USED ON (parent drawings if any)"],
           (26, 40, 60), rows, tab="green", empty="No root assemblies in this run",
           note=("About this sheet",
                 "Top-level assemblies of the tree. USED ON files belong to parent "
                 "assemblies - download them from Teamcenter if not already in the tree "
                 "(check the Files sheet).", "green"))


def _mismatch_sheet(mism_ws, ctx, names):
    mism_data = ctx.get("mismatches", [])  # [(parent, child, actual_list|str), ...]
    bug_data = ctx.get("used_on_bugs", [])  # [(parent, child), ...]
    rows = []
    for pstem, cstem, actual in mism_data:
        if isinstance(actual, (list, tuple)):
            actual = ", ".join(actual) if actual else "none"
        rows.append([pstem, names.get(pstem, ""), cstem, names.get(cstem, ""), actual or "-"])
    ncols = _table(mism_ws, ["Parent drawing", "Parent name", "Child part", "Child name",
                             "Child USED ON (actual)"], (26, 40, 26, 40, 44), rows,
                   tab="amber", empty="No BOM / USED ON mismatches",
                   note=("About this sheet",
                         "Top table: the parent's BOM lists a part but that part's USED ON "
                         "omits the parent - verify the USED ON number.\n\nLower table "
                         "(USED ON BUGS): a child's USED ON names a parent whose BOM does "
                         "not list it. BOM is the only source of truth, so these are data "
                         "bugs, not placement edges.", "amber"))
    base = len(mism_data) + 6
    sec = mism_ws.cell(row=base, column=1, value=f"USED ON BUGS ({len(bug_data)})")
    sec.font = _font(12, True, DARK)
    mism_ws.row_dimensions[base].height = 26
    mism_ws.cell(row=base + 1, column=1,
                 value="Child USED ON lists a parent that the parent's BOM does not list "
                       "- fix the USED ON number.").font = _font(9, False, MUTED, italic=True)
    sub = base + 2
    for ci, text in enumerate(("Parent drawing", "Parent name", "Child part", "Child name"),
                              start=1):
        c = mism_ws.cell(row=sub, column=ci, value=text)
        c.font = _font(10, True, WHITE)
        c.fill = _solid(NAVY)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    mism_ws.row_dimensions[sub].height = 24
    if not bug_data:
        mism_ws.merge_cells(start_row=sub + 1, start_column=1, end_row=sub + 1, end_column=4)
        c = mism_ws.cell(row=sub + 1, column=1, value="\u2713  No USED ON bugs")
        c.font = _font(10, True, PAL["green"]["text"])
        c.fill = _solid(PAL["green"]["fill"])
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, (pstem, cstem) in enumerate(bug_data):
        r = sub + 1 + i
        vals = (pstem, names.get(pstem, ""), cstem, names.get(cstem, ""))
        mism_ws.row_dimensions[r].height = ROW_H
        for ci, v in enumerate(vals, start=1):
            c = mism_ws.cell(row=r, column=ci, value=v)
            c.fill = _solid(STRIPE if i % 2 else WHITE)
            c.border = ROW_BORDER
            c.font = _font(10, ci == 1, INK)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return ncols


def _titleblock_sheet(tb_ws, ctx, names):
    data = ctx.get("titleblock_mismatches", [])  # [(stem, field, filename, tb)]
    rows = [[stem, names.get(stem, "") or "-", field, fval, tval]
            for stem, field, fval, tval in sorted(data)]
    _table(tb_ws, ["Stem (current file)", "Name", "Field", "In filename", "In title block"],
           (30, 34, 12, 26, 26), rows, tab="amber", pills={3: "amber"},
           aligns=["left", "left", "center", "left", "left"],
           empty="Every filename agrees with its title block",
           note=("About this sheet",
                 "The filename disagrees with the drawing's own title block - verify the "
                 "filename or re-download from Teamcenter.", "amber"))


def _files_sheet(files, ctx, names, live):
    chk_set = set(ctx.get("chk", []))
    out = ctx.get("output", "")
    rows, paths = [], []

    def add(p, status):
        stem = canonical_stem(p.stem) or p.stem.upper()
        parts = stem.split("_")
        rev = parts[1] if len(parts) > 1 and parts[1] else "-"
        rows.append([stem, names.get(stem, "") or "-", rev, status,
                     str(Path(_rel(p, out)).parent)])
        paths.append(p)

    for p in sorted(live["tree"], key=lambda q: str(q).lower()):
        stem = canonical_stem(p.stem) or p.stem.upper()
        add(p, "CHK (unapproved)" if stem in chk_set else "Active")
    for p in sorted(live["orph"], key=lambda q: str(q).lower()):
        add(p, "ORPHAN (parked)")
    pal = {"Active": "green", "CHK (unapproved)": "amber", "ORPHAN (parked)": "orange"}
    _table(files, ["Drawing (stem)", "Name", "Revision", "Status", "Folder (relative to output)"],
           (26, 40, 10, 20, 60), rows, tab="blue",
           aligns=["left", "left", "center", "center", "left"],
           pills={4: lambda row: pal[row[3]]},
           link_first=lambda i, row: _file_uri(paths[i]),
           empty="No PDFs in the output tree yet", max_width=110,
           note=("About this sheet",
                 "Live scan of the output folder. Click a drawing to open its PDF; use "
                 "the filter arrows to slice by status or revision.", "blue"))


def _history_sheet(hist, ctx, counters, missing_pairs, path, log):
    entry = [str(ctx.get("run_time", "")), str(ctx.get("run_mode", "")),
             counters.get("scanned", 0), counters.get("copies", 0),
             len(missing_pairs), len(set(ctx.get("chk", []))),
             len(ctx.get("orphans", [])), ctx.get("report_txt", "")]
    hist_rows = _load_history(path, log) + [entry]
    _save_history(path, hist_rows, log)
    rows = [list(r) for r in sorted(hist_rows[-HISTORY_MAX_RUNS:], reverse=True)]
    n = len(rows)
    _table(hist, ["Run time", "Mode", "PDFs scanned", "Copies", "Missing", "CHK",
                  "Orphans", "Report (txt)"], (20, 16, 14, 11, 11, 9, 11, 46), rows,
           tab="blue",
           aligns=["left", "center", "center", "center", "center", "center", "center", "left"],
           empty="No runs recorded yet",
           note=("About this sheet",
                 "One row per run, newest first (last 200 kept in a sidecar .history.json "
                 "next to this workbook). The chart shows how the open items trend.", "blue"))
    if n:
        grey = PAL["grey"]
        green = PAL["green"]
        hist.conditional_formatting.add(
            f"B2:B{n + 1}", FormulaRule(
                formula=['ISNUMBER(SEARCH("DRY",B2))'],
                fill=PatternFill(start_color=grey["fill"], end_color=grey["fill"], fill_type="solid"),
                font=Font(name=FONT, bold=True, color=grey["text"])))
        hist.conditional_formatting.add(
            f"B2:B{n + 1}", FormulaRule(
                formula=['ISNUMBER(SEARCH("EXEC",B2))'],
                fill=PatternFill(start_color=green["fill"], end_color=green["fill"], fill_type="solid"),
                font=Font(name=FONT, bold=True, color=green["text"])))
        hist.conditional_formatting.add(
            f"D2:D{n + 1}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                        color="93C5FD", showValue=True))
    if n >= 2:
        lc = LineChart()
        lc.title = "Open items per run"
        lc.height, lc.width = 7.5, 18
        lc.y_axis.majorGridlines = None
        lc.x_axis.delete = False
        lc.y_axis.delete = False
        lc.x_axis.scaling.orientation = "maxMin"   # rows are newest-first
        lc.y_axis.crosses = "max"                  # keep the value axis on the left
        for col, pal in ((5, "red"), (6, "amber"), (7, "orange")):
            lc.add_data(Reference(hist, min_col=col, min_row=1, max_row=min(n, 30) + 1),
                        titles_from_data=True)
            s = lc.series[-1]
            s.graphicalProperties.line.solidFill = PAL[pal]["accent"]
            s.graphicalProperties.line.width = 22000
            s.smooth = False
        lc.set_categories(Reference(hist, min_col=1, min_row=2, max_row=min(n, 30) + 1))
        hist.add_chart(lc, "J11")


def refresh_from_run(output, run_mode, run_time, counters, missing, chk,
                     orphans, report_txt, names=None, roots=None, mismatches=None,
                     used_on_bugs=None, titleblock_mismatches=None, scanned=None,
                     watermarks=None, log=None):
    """Called by the organizer after every run (execute AND dry-run).

    Dry-runs into a fresh --output never create the folder; the workbook is
    still produced, so it creates the output folder if missing (report-only).
    """
    outdir = Path(output)
    outdir.mkdir(parents=True, exist_ok=True)
    xlsx = os.path.join(str(outdir), "organize_fermi_report.xlsx")
    return build_workbook(xlsx, {
        "output": str(output),
        "run_mode": run_mode,
        "run_time": run_time,
        "counters": counters,
        "missing": missing,
        "chk": chk,
        "orphans": orphans,
        "roots": roots or [],
        "mismatches": mismatches or [],
        "used_on_bugs": used_on_bugs or [],
        "titleblock_mismatches": titleblock_mismatches or [],
        "scanned": scanned or [],
        "watermarks": watermarks or [],
        "report_txt": os.path.basename(report_txt or ""),
        "names": names or {},
    }, log=log)
